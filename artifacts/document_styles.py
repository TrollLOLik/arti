"""Print-friendly document exports, with no changes to source values or evidence.

The report and workbook share the renderer's semantic palette.  Workbook cell
coordinates stay source-relative; the print header supplies the document title
without inserting rows into the dataset.  All exported cell content is literal.
"""
from dataclasses import asdict
from io import BytesIO
from math import ceil

from artifacts.styles import semantic_theme
from materials.types import MaterialError, canonical


FONT = 'DejaVu Sans'
STATUS_LABELS = {
    'observed': 'Наблюдение', 'confirmed': 'Подтверждено',
    'proposed': 'Предложение', 'unknown': 'Неизвестно', 'fiction': 'Вымысел',
}


def _hex(theme, role):
    return theme[role].lstrip('#')


def _quantity(element):
    quantity = element.get('quantity')
    if not quantity:
        return 'Значение не задано'
    text = quantity['value'] + ' ' + quantity['unit']
    if 'lower' in quantity or 'upper' in quantity:
        text += '\nДиапазон: {} … {} {}'.format(
            quantity.get('lower', quantity['value']),
            quantity.get('upper', quantity['value']), quantity['unit'])
    return text


def _status(element):
    value = element['status']
    return STATUS_LABELS.get(value, value) + ' (' + value + ')'


def _set_font(style, size, color, bold=False):
    from docx.shared import Pt, RGBColor
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def _word_field(paragraph, name):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), name)
    paragraph._p.append(field)


def _word_table(doc, headings, records, theme, widths):
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(table.rows[0].cells, headings):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in records:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for key, value in (('val', 'single'), ('sz', '4'), ('color', _hex(theme, 'border'))):
            border.set(qn('w:' + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), _hex(theme, 'heading' if index == 0 else 'soft_surface' if index % 2 else 'surface'))
            props.append(shading)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '100')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.12
                paragraph.paragraph_format.keep_with_next = index == 0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.font.bold = index == 0
                    run.font.color.rgb = RGBColor.from_string(_hex(theme, 'inverse' if index == 0 else 'text'))
    return table


def report_docx(spec, revision, files):
    """Render an already authorized spec; authorization stays with the caller."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    # Paper exports use the shared print-safe theme, including for night-mode
    # artifacts: office viewers do not reliably print document page backgrounds.
    theme = semantic_theme()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.3)
    normal = doc.styles['Normal']
    _set_font(normal, 10.5, _hex(theme, 'text'))
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'ru-RU')
    normal.element.get_or_add_rPr().append(language)
    for name, size, role in (('Title', 24, 'text'), ('Heading 1', 15, 'heading'), ('Heading 2', 12, 'heading')):
        style = doc.styles[name]
        _set_font(style, size, _hex(theme, role), bold=True)
        style.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        # The bundled Word template may carry a blue title rule. Remove that
        # inherited decoration instead of letting theme colors escape the palette.
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    for name, size in (('Arti Metadata', 9), ('Arti Source', 8.5), ('Arti Caption', 9)):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        _set_font(style, size, _hex(theme, 'secondary'))
        style.paragraph_format.space_after = Pt(6)
    doc.core_properties.title = spec['title']
    doc.core_properties.subject = 'Отчёт по спецификации Arti'
    doc.core_properties.author = 'Arti'
    header = section.header.paragraphs[0]
    header.text = 'ARTI  /  ОТЧЁТ'
    header.style = doc.styles['Arti Metadata']
    footer = section.footer.paragraphs[0]
    footer.style = doc.styles['Arti Metadata']
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Arti · Редакция {} · Страница '.format(revision))
    _word_field(footer, 'PAGE')
    footer.add_run(' из ')
    _word_field(footer, 'NUMPAGES')

    doc.add_paragraph(spec['title'], 'Title')
    doc.add_paragraph('Редакция {} · Формат {} · Элементов {}'.format(
        revision, spec['format'], len(spec['elements'])), 'Arti Metadata')
    doc.add_heading('Сводка', 1)
    _word_table(doc, ['Элемент', 'Статус', 'Значение'], [
        [e['label'], _status(e), _quantity(e)] for e in spec['elements']
    ], theme, [3.05, 1.7, 2.25])

    doc.add_heading('Подробности', 1)
    for element in spec['elements']:
        doc.add_heading(element['label'], 2)
        doc.add_paragraph(_status(element) + ' · ID ' + element['id'], 'Arti Metadata')
        if element.get('text'):
            doc.add_paragraph(element['text'])
        doc.add_paragraph(_quantity(element))
        if 'when' in element:
            doc.add_paragraph('Дата: ' + element['when'])
    if spec.get('relations'):
        doc.add_heading('Связи', 1)
        labels = {e['id']: e['label'] for e in spec['elements']}
        for relation in spec['relations']:
            doc.add_paragraph('{} → {} · {}'.format(
                labels[relation['from']], labels[relation['to']], relation['kind']))
            if relation.get('label'):
                doc.add_paragraph(relation['label'])
            # Full record retains relation IDs, kind and proof without inference.
            doc.add_paragraph(canonical(relation), 'Arti Source')
    if spec.get('axis'):
        doc.add_heading('Шкала и единицы', 1)
        doc.add_paragraph(canonical(spec['axis']), 'Arti Source')
    if spec.get('questions'):
        doc.add_heading('Открытые вопросы', 1)
        for question in spec['questions']:
            doc.add_paragraph(question)

    doc.add_heading('Источники', 1)
    for element in spec['elements']:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(element['label'] + ' · ' + element['id']).bold = True
        if element.get('proof'):
            doc.add_paragraph(canonical(element['proof']), 'Arti Source')
        else:
            doc.add_paragraph('Источник не указан. Статус: ' + _status(element), 'Arti Source')

    pngs = [(name, pixels) for name, pixels in files.items() if name.endswith('.png')]
    for index, (name, pixels) in enumerate(pngs):
        doc.add_page_break()
        doc.add_heading('Визуальный результат' if index == 0 else 'Визуальный результат продолжение', 1)
        caption = ('Художественная иллюстрация; не свидетельство фактов.'
                   if name.startswith('illustration-') else 'Программный рендер проверяемой спецификации.')
        paragraph = doc.add_paragraph(caption, 'Arti Caption')
        paragraph.paragraph_format.keep_with_next = True
        # Fit both axes, including unusually tall illustrations, without crop.
        from PIL import Image
        with Image.open(BytesIO(pixels)) as image:
            width = min(6.35, 7.85 * image.width / image.height)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(BytesIO(pixels), width=Inches(width))
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _literal_cell(sheet, row, column, value):
    """Exact strings protect all sheets from formulas and float round-trips."""
    text = '' if value is None else str(value)
    # openpyxl silently slices strings beyond Excel's cell limit. Reject first,
    # including provenance: a pretty export must never hide lost source text.
    if len(text) > 32767:
        raise MaterialError('document_cell_text_budget')
    cell = sheet.cell(row, column, text)
    cell.data_type = 's'
    cell.number_format = '@'
    return cell


def _style_sheet(sheet, theme, title, header_row, widths=None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.sheet_properties.tabColor = _hex(theme, 'heading')
    # A title in the print header preserves original data addresses for readers
    # tracing Sources back to the unshifted matrix. Escape Excel header commands.
    sheet.oddHeader.center.text = title.replace('&', '&&')
    sheet.oddHeader.center.font = FONT + ',Bold'
    sheet.oddHeader.center.size = 14
    sheet.oddHeader.center.color = _hex(theme, 'heading')
    sheet.oddFooter.left.text = 'Arti · ' + sheet.title
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.left.color = _hex(theme, 'secondary')
    sheet.oddFooter.right.text = 'Страница &P из &N'
    sheet.oddFooter.right.size = 8
    sheet.page_setup.orientation = 'landscape' if sheet.max_column > 4 else 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.7, bottom=.5, header=.2, footer=.2)
    sheet.print_options.horizontalCentered = True
    sheet.print_area = 'A1:{}{}'.format(get_column_letter(sheet.max_column), sheet.max_row)
    sheet.freeze_panes = 'A{}'.format(header_row + 1) if header_row else None
    if header_row:
        sheet.print_title_rows = '{}:{}'.format(header_row, header_row)
        sheet.auto_filter.ref = 'A{}:{}{}'.format(header_row, get_column_letter(sheet.max_column), sheet.max_row)
    for col in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, col).value or '') for row in range(1, sheet.max_row + 1)]
        width = widths[col - 1] if widths else min(48, max(14, max((len(v) for v in values), default=0) + 3))
        sheet.column_dimensions[get_column_letter(col)].width = width
    total_width = sum(sheet.column_dimensions[get_column_letter(col)].width for col in range(1, sheet.max_column + 1))
    if total_width > 95:
        sheet.page_setup.orientation = 'landscape'
    if total_width > 150:
        # Very wide source matrices need horizontal pages, not microscopic text.
        sheet.sheet_properties.pageSetUpPr.fitToPage = False
        sheet.page_setup.fitToWidth = 0
        sheet.page_setup.scale = 100
        sheet.print_title_cols = 'A:A'
    thin = Side(style='thin', color=_hex(theme, 'border'))
    for row in sheet:
        height = 26
        for cell in row:
            is_header = cell.row == header_row
            cell.font = Font(name=FONT, size=10, bold=is_header, color=_hex(theme, 'inverse' if is_header else 'text'))
            cell.fill = PatternFill('solid', fgColor=_hex(theme, 'heading' if is_header else 'soft_surface' if cell.row % 2 == 0 else 'surface'))
            cell.border = Border(bottom=thin, right=thin)
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
            width = sheet.column_dimensions[cell.column_letter].width
            lines = sum(max(1, ceil(len(line) / max(1, width - 4))) for line in str(cell.value or '').split('\n'))
            height = max(height, lines * 16 + 10)
            if lines * 16 + 10 > 409:
                from openpyxl.comments import Comment
                prior = cell.comment.text + '\n' if cell.comment else ''
                cell.comment = Comment(prior + 'Длинный текст может быть обрезан при печати. Полное значение сохранено в ячейке.', 'Arti')
        # Excel's own maximum row height is 409 points. No cell values are cut.
        sheet.row_dimensions[row[0].row].height = min(409, height)


def dataset_xlsx(data, matrix):
    """Style the native snapshot without changing data, units, or provenance."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    theme = semantic_theme()
    book = Workbook()
    book.properties.creator = 'Arti'
    book.properties.title = data.name
    book.properties.subject = 'Данные и исходные свидетельства'
    sheet = book.active
    sheet.title = 'Verified data'
    for row, values in enumerate(matrix, 1):
        for col, value in enumerate(values, 1):
            _literal_cell(sheet, row, col, value)
    sources = book.create_sheet('Sources')
    headings = ['Ячейка', 'Оригинал', 'Нормализованное', 'Единица', 'Источник']
    for col, value in enumerate(headings, 1):
        _literal_cell(sources, 1, col, value)
    for row, item in enumerate(data.cells, 2):
        values = [item.address, item.raw, item.normalized.value, item.normalized.unit, canonical(asdict(item.source))]
        for col, value in enumerate(values, 1):
            _literal_cell(sources, row, col, value)
        details = 'Нормализация: ' + canonical(asdict(item.normalized)) + '\nИсточник: ' + canonical(asdict(item.source))
        if item.formula is not None:
            details += '\nИсходная формула (не выполняется): ' + item.formula
        if item.cached_raw is not None:
            details += '\nИсходное кэшированное значение: ' + str(item.cached_raw)
        sheet.cell(item.row + 1, item.column + 1).comment = Comment(details, 'Arti')
        sources.cell(row, 3).comment = Comment(details, 'Arti')
    header = data.policy.header_row + 1 if data.policy.header_row is not None else None
    if header and header > sheet.max_row:
        header = None
    _style_sheet(sheet, theme, data.name or 'Данные', header)
    _style_sheet(sources, theme, 'Источники и исходные значения', 1, [10, 25, 23, 12, 72])
    # Explicit types and missing-value reasons are available without changing
    # the displayed source cell; a missing value is never displayed as zero.
    for item in data.cells:
        cell = sheet.cell(item.row + 1, item.column + 1)
        if item.normalized.kind == 'number' and item.row + 1 != header:
            from openpyxl.styles import Alignment
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True, indent=1)
    sheet['A1'].comment = Comment((sheet['A1'].comment.text + '\n' if sheet['A1'].comment else '') +
        'Набор: ' + data.name + '\nПокрытие: ' + data.coverage + '\nОграничения: ' + canonical(data.limitations) +
        '\nТочные значения сохранены текстом. Тип, единица и диапазон указаны в примечаниях; источники на листе Sources.', 'Arti')
    out = BytesIO()
    book.save(out)
    return out.getvalue()
