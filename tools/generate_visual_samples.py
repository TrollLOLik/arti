"""Build owned, source-validated visual examples with disposable PostgreSQL.

This corpus is synthetic, explicitly labelled, and never reads real documents.
Use the same offline guard and disposable database policy as regression tests.
"""
import argparse
import asyncio
import base64
from copy import deepcopy
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile

from tools.verify_native_agent_scenarios import offline_guard


def source_documents():
    from docx import Document
    from openpyxl import Workbook
    paragraphs = [
        'СИНТЕТИЧЕСКИЙ ПРИМЕР. Проект «Северный сад» вымышлен; это контроль оформления, а не пользовательские сведения.',
        'Задача выставки — показать посетителям, как городской сад меняется в течение года.',
        'В экспозиции предложены три раздела: почва, вода и сезонные растения. Решение о составе ещё не принято.',
        'Сведения о доступности площадки и окончательном бюджете в исходном документе отсутствуют.',
        '2026-10-06 — подготовить перечень источников; это предложение, а не выполненная работа.',
        '2026-10-12 — проверить макет подписей и таблиц; это предложение, а не утверждённый срок.',
        '2026-10-20 — обсудить результаты проверки; решение о запуске ещё не принято.',
    ]
    doc = Document(); doc.add_heading('Северный сад / исходный материал', 0)
    for line in paragraphs: doc.add_paragraph(line)
    out = BytesIO(); doc.save(out)
    book = Workbook(); ws = book.active; ws.title = 'Показатели'
    for row in [
        ['Показатель', 'Изменение, RUB'],
        ['Изменение материалов', '-1200'],
        ['Изменение монтажа', '2400'],
        ['Малое изменение доставки', '0.000001'],
        ['Неизвестная стоимость площадки', 'нет данных'],
    ]: ws.append(row)
    xout = BytesIO(); book.save(xout)
    return paragraphs, out.getvalue(), xout.getvalue()


async def generate(target):
    from tests.support.database import isolated_database
    from cognition.repositories import ensure_schema
    from materials.repository import MaterialRepository
    from materials.service import MaterialService
    from materials.storage import LocalBlobStore
    from materials.types import AccessContext, MaterialScope, EvidenceRef
    from materials.datasets import DatasetPolicy, ColumnPolicy
    from projects.repository import ProjectRepository
    from artifacts.revisions import ArtifactRepository
    from artifacts.spec import ArtifactSpec, FORMATS
    from artifacts.export import export_current
    from artifacts.validation import validate_evidence
    from agents.tools.core import build_registry
    from agents.tools.registry import ToolContext
    from pypdf import PdfReader
    from PIL import Image
    import pypdfium2 as pdfium

    target.mkdir(parents=True, exist_ok=True)
    paragraphs, docx, xlsx = source_documents()
    (target / 'source-synthetic.docx').write_bytes(docx)
    (target / 'source-synthetic.xlsx').write_bytes(xlsx)
    cases = []
    with tempfile.TemporaryDirectory() as storage:
        async with isolated_database() as pool:
            await ensure_schema(pool)
            materials = MaterialRepository(pool)
            service = MaterialService(materials, LocalBlobStore(storage))
            actor = AccessContext(MaterialScope('arti', 9191, -1, 'private'), 7171, 'user:7171')
            project = await ProjectRepository(materials).create(actor, 'Синтетическая проверка оформления')
            document = await service.ingest(docx, 'source-synthetic.docx', actor, 'synthetic:document', 'synthetic:document')
            extraction, bundle = await service.extract(document['id'], actor)
            workbook = await service.ingest(xlsx, 'source-synthetic.xlsx', actor, 'synthetic:workbook', 'synthetic:workbook')
            datasets = await service.datasets(workbook['id'], actor, policy=DatasetPolicy(columns=(ColumnPolicy(1, unit='RUB', locale='en'),)))
            dataset = datasets[0]
            def quote(index, label, **kwargs):
                text = paragraphs[index]
                block = next(b for b in bundle.blocks if b.text == text)
                ref = EvidenceRef(document['id'], bundle.asset_version, extraction, block.block_id, block.locator)
                return dict(id=f'e{index}', label=label, text=text, status='observed',
                            proof=dict(kind='quote', source=asdict(ref), quote=text), **kwargs)
            def spec(title, format, elements, **kwargs):
                return dict(contract='artifact-1', title=title, format=format, elements=elements, relations=[],
                            questions=['Какие сведения о площадке и бюджете нужно уточнить?'], **kwargs)
            cards = spec('Северный сад / учебный пример', 'cards', [
                quote(1, 'Назначение'), quote(2, 'Предложенная структура'),
                quote(3, 'Чего пока нет'),
            ])
            numbers = []
            for i, address in enumerate(('B2', 'B3', 'B4')):
                cell = dataset.cell(address); label = dataset.cell(f'A{i+2}').raw
                assert cell.normalized.kind == 'number'
                numbers.append(dict(id=f'n{i}', label=label, status='observed',
                    quantity=dict(value=cell.normalized.value, unit=cell.normalized.unit),
                    proof=dict(kind='dataset', dataset_id=dataset.id, address=address)))
            numbers.append(dict(id='missing', label='Стоимость площадки', status='unknown', text='В источнике: «нет данных». Значение не заменено нулём.'))
            statistical = spec('Изменения бюджета / учебный пример', 'statistical', numbers, axis=dict(scale='linear', unit='RUB'))
            timeline = spec('Учебный пример / предложенные этапы', 'timeline', [
                quote(4, 'Подготовка источников', order=1, when='2026-10-06'),
                quote(5, 'Проверка макета', order=2, when='2026-10-12'),
                quote(6, 'Обсуждение результата', order=3, when='2026-10-20'),
            ])
            timeline['relations'] = [dict(id='r1', kind='sequence', **{'from':'e4','to':'e5'}), dict(id='r2', kind='sequence', **{'from':'e5','to':'e6'})]
            source = asdict(EvidenceRef(document['id'], bundle.asset_version, extraction, bundle.blocks[0].block_id, bundle.blocks[0].locator))
            repo = ArtifactRepository(materials); repo.service = service
            rows = {}
            for name, value in [('01-document-summary', cards), ('02-statistics', statistical), ('03-timeline', timeline)]:
                row = await repo.create(project.id, actor, value, sources=[source]); rows[name] = row
                files, _ = await export_current(repo, row['id'], actor, revision=row['revision'])
                folder = target / name; folder.mkdir(exist_ok=True)
                for filename, content in files.items(): (folder / filename).write_bytes(content)
                pages = PdfReader(BytesIO(files['report.pdf'])).pages
                compact = lambda text: ''.join(text.split())
                extracted = compact(' '.join(page.extract_text() for page in pages))
                exact = all(compact(e['label']) in extracted and compact(e.get('text','')) in extracted
                            and compact(e.get('quantity',{}).get('value','')) in extracted for e in value['elements'])
                assert exact, name
                document_pdf = pdfium.PdfDocument(files['report.pdf'])
                for page_index in range(len(document_pdf)):
                    rendered = document_pdf[page_index].render(scale=.5).to_pil()
                    rendered.save(folder / f'pdf-page-{page_index+1}.png')
                document_pdf.close()
                with Image.open(BytesIO(files['page-1.png'])) as picture:
                    picture.resize((390,520)).save(folder / 'phone-preview.png')
                cases.append(dict(name=name, format=value['format'], pages=len(pages), exact_content_preserved=exact,
                    source_checks=row['checks'], files={filename:sha256(data).hexdigest() for filename,data in files.items()},
                    pixel_review='pending_separate_inspection'))
            registry = build_registry()
            async def guard(): await repo.get(rows['01-document-summary']['id'], actor)
            context = ToolContext(actor, service, project.id, 'synthetic', 'synthetic', guard)
            svg_tool_sizes = {}
            # Every declared format exercises real evidence and render validation.
            for format in sorted(FORMATS):
                value = deepcopy(statistical if format in ('comparison','statistical','table') else timeline if format in ('timeline','roadmap') else cards)
                value['format'] = format
                if format != 'statistical': value.pop('axis', None)
                row = await repo.create(project.id, actor, value, sources=[source])
                files, _ = await export_current(repo, row['id'], actor, revision=1)
                folder = target / 'all-formats'; folder.mkdir(exist_ok=True)
                (folder / f'{format}.png').write_bytes(files['page-1.png'])
                svg_result = await registry.call('artifact.export', dict(id=row['id'], revision=1, format='svg'), context, version='1')
                svg_tool_sizes[format] = sum(len(item['base64']) for item in svg_result.outputs['files'])
            for tool, arguments in [
                ('documents.report', dict(artifact_id=rows['01-document-summary']['id'], revision=1, format='docx')),
                ('documents.data_export', dict(dataset_id=dataset.id, format='xlsx')),
            ]:
                result = await registry.call(tool, arguments, context, version='1')
                for item in result.outputs['files']: (target / item['name']).write_bytes(base64.b64decode(item['base64'], validate=True))
            from artifacts.material_cards import extraction_card, dataset_card, render_card
            from materials.dataset_quality import diagnose
            (target / 'extraction-quality.png').write_bytes(render_card(extraction_card(bundle,'source-synthetic.docx')))
            (target / 'dataset-quality.png').write_bytes(render_card(dataset_card(datasets,[diagnose(dataset)])))
    report = dict(schema='arti-source-bound-visual-samples-1', synthetic=True,
        scope='owned DOCX/XLSX ingested, extracted, source-validated and rendered offline', cases=cases,
        all_nine_formats_rendered=True, svg_tool_base64_bytes=svg_tool_sizes, live_model=False, live_telegram=False, production_data=False)
    (target / 'manifest.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(dict(output=str(target),cases=len(cases),all_nine_formats_rendered=True),ensure_ascii=False))


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--output',default='temp/visual-design-samples')
    args=parser.parse_args(); offline_guard(); asyncio.run(generate(Path(args.output)))

if __name__ == '__main__': main()
