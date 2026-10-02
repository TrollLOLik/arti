"""Original Russian/English fixtures; no user documents or external downloads."""
from io import BytesIO
from pathlib import Path
from functools import lru_cache
from PIL import Image, ImageDraw, ImageFont


def font_path():
    for path in (Path('C:/Windows/Fonts/arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')):
        if path.is_file(): return str(path)
    raise RuntimeError('Cyrillic fixture font missing')


def scan(*,angle=0,heldout=False,table=False):
    image=Image.new('RGB',(1200,1500),'white'); draw=ImageDraw.Draw(image)
    font=ImageFont.truetype(font_path(),38)
    text=['Отчёт проекта Арти','Дата: 30 сентября 2026','Аренда: 1200 рублей','Доставка: 300 рублей','Итого: 1500 рублей',
           'Решение ещё не принято.','Проверка источников обязательна.','Документ содержит наблюдения,','а не разрешение на действия.']
    if heldout:
        text=['Контрольная ведомость','Дата: 18 ноября 2026','Материалы: 2750 рублей','Монтаж: 840 рублей','Итого: 3590 рублей',
               'Предложение отменено.','Это другая контрольная выборка.','Показатели требуют проверки.']
    for i,line in enumerate(text): draw.text((100,90+i*90),line,font=font,fill='black')
    if table:
        for x in (100,640,1080): draw.line((x,960,x,1320),fill='black',width=3)
        for y in (960,1080,1200,1320): draw.line((100,y,1080,y),fill='black',width=3)
        for ri,row in enumerate([['Статья','Рубли'],['Аренда','1200'],['Доставка','300']]):
            for ci,value in enumerate(row): draw.text((120+ci*540,990+ri*120),value,font=font,fill='black')
    return image.rotate(angle,expand=True,fillcolor='white') if angle else image


def scanned_pdf(**kwargs):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    output=BytesIO(); image=scan(**kwargs)
    width,height=image.width/2,image.height/2
    c=canvas.Canvas(output,pagesize=(width,height)); c.drawImage(ImageReader(image),0,0,width=width,height=height); c.save()
    return output.getvalue()


@lru_cache(maxsize=1)
def structured_pdf():
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    pdfmetrics.registerFont(TTFont('FixtureRussian',font_path()))
    output=BytesIO(); c=canvas.Canvas(output,pagesize=(600,800))
    for page in (1,2):
        c.setFont('FixtureRussian',10); c.drawString(50,770,'Отчёт Арти'); c.drawString(280,25,f'Страница {page}')
        c.setFont('FixtureRussian',18); c.drawString(50,725,'Результаты анализа')
        c.setFont('FixtureRussian',12)
        if page==1:
            for i,value in enumerate(['Левая колонка 1','Левая колонка 2','Левая колонка 3']): c.drawString(50,675-i*27,value)
            for i,value in enumerate(['Правая колонка 1','Правая колонка 2','Правая колонка 3']): c.drawString(335,675-i*27,value)
            c.drawString(50,570,'1. Проверить источники')
            img=Image.new('RGB',(250,140),'white'); draw=ImageDraw.Draw(img)
            draw.rectangle((20,20,75,120),fill='blue'); draw.rectangle((120,60,175,120),fill='red')
            c.drawImage(ImageReader(img),50,330,width=250,height=140)
            c.drawString(50,307,'Рис. 1. Сравнение показателей')
            start,rows=220,[['Статья','Рубли'],['Аренда','1200'],['Доставка','300']]
        else:
            start,rows=690,[['Статья','Рубли'],['Монтаж','840'],['Материалы','2750']]
            image=scan(heldout=True).crop((70,230,1080,590))
            c.drawImage(ImageReader(image),50,270,width=500,height=180)
            c.drawString(50,247,'Рис. 2. Фрагмент контрольной ведомости')
        for x in (50,330,550): c.line(x,start,x,start-120)
        for y in (start,start-40,start-80,start-120): c.line(50,y,550,y)
        for ri,row in enumerate(rows):
            for ci,value in enumerate(row): c.drawString(60+ci*280,start-27-ri*40,value)
        c.showPage()
    c.save(); return output.getvalue()


def rich_docx():
    from docx import Document
    from tests.materials.fixtures import png
    document=Document(); document.sections[0].header.paragraphs[0].text='Служебный заголовок'
    document.sections[0].footer.paragraphs[0].text='Подвал документа'
    document.add_heading('Бюджет проекта',level=1)
    document.add_paragraph('Проверить числа',style='List Bullet')
    table=document.add_table(rows=3,cols=3)
    table.cell(0,0).merge(table.cell(0,2)).text='Общий бюджет'
    table.cell(1,0).merge(table.cell(2,0)).text='Работы'
    table.cell(1,1).text='Монтаж'; table.cell(1,2).text='840'
    table.cell(2,1).text='Материалы'; table.cell(2,2).text='2750'
    nested=table.cell(2,1).add_table(rows=1,cols=2); nested.cell(0,0).text='Код'; nested.cell(0,1).text='42'
    document.add_picture(BytesIO(png()))
    document.add_paragraph('Рис. 1. Контрольная схема',style='Caption')
    output=BytesIO(); document.save(output); return output.getvalue()


def damaged_page_pdf():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import EncodedStreamObject,NameObject
    writer=PdfWriter(); writer.append_pages_from_reader(PdfReader(BytesIO(structured_pdf())))
    broken=EncodedStreamObject(); broken._data=b'broken flate stream'; broken[NameObject('/Filter')]=NameObject('/FlateDecode')
    writer.pages[0][NameObject('/Contents')]=writer._add_object(broken)
    out=BytesIO(); writer.write(out); return out.getvalue()


def conflicting_text_layer_pdf():
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont('FixtureRussian',font_path()))
    out=BytesIO(); c=canvas.Canvas(out,pagesize=(600,750))
    c.drawImage(ImageReader(scan()),0,0,width=600,height=750)
    hidden=c.beginText(50,600); hidden.setFont('FixtureRussian',19); hidden.setTextRenderMode(3)
    hidden.textLine('Аренда: 9999 рублей'); c.drawText(hidden); c.save()
    return out.getvalue()
