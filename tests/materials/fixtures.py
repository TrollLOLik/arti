"""Synthetic, redistributable corpus with known evidence and exact numbers."""
from io import BytesIO
from PIL import Image, ImageDraw


def png():
    output = BytesIO()
    image = Image.new('RGB', (240, 120), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 70, 100), fill='blue')
    draw.rectangle((100, 50, 160, 100), fill='red')
    image.save(output, format='PNG')
    return output.getvalue()


def docx():
    from docx import Document
    document = Document()
    document.add_paragraph('Бюджет встречи')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Статья'
    table.cell(0, 1).text = 'Рубли'
    table.cell(1, 0).text = 'Аренда'
    table.cell(1, 1).text = '1200'
    document.add_paragraph('Предложение, ещё не принято: дата 12 октября.')
    document.add_picture(BytesIO(png()))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for number in (1200, 300):
        page = writer.add_blank_page(300, 200)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 16 Tf 30 100 Td (Budget {number}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()
