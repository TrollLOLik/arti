"""Offline office-export regression tests and reproducible synthetic samples.

Generate samples with:
  python -m tests.materials.test_document_visuals --samples /tmp/arti-office
"""
import base64
from copy import deepcopy
from dataclasses import asdict, replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

from artifacts.document_styles import dataset_xlsx, report_docx
from artifacts.spec import ArtifactSpec
from materials.datasets import DataCell, Dataset, DatasetPolicy, Value
from materials.types import EvidenceRef, Locator, MaterialError, canonical


def synthetic_dataset():
    values = [
        ['Показатель', 'Сумма RUB', 'Комментарий'],
        ['Корректировка', '-0.000001', 'Точное отрицательное значение'],
        ['Наблюдение', '2.345', 'Диапазон 2.300 … 2.400'],
        ['Нулевое значение', '0', 'Ноль указан в источнике'],
        ['Нет наблюдения', 'нет данных', 'Отсутствие не равно нулю'],
        ['Текст источника', '=HYPERLINK("https://example.invalid","пример")', '+Текст'],
        ['Проверка литералов', '@SUM(1,2)', '-Текст'],
    ]
    cells = []
    for row, items in enumerate(values):
        for column, raw in enumerate(items):
            address = chr(65 + column) + str(row + 1)
            source = EvidenceRef('synthetic-ledger', 1, 'synthetic-extraction', 'cell-' + address,
                                 Locator('cell', sheet='Пример', cell=address))
            if column == 1 and row in (1, 2, 3):
                normalized = Value('number', raw, 'RUB', '2.300' if row == 2 else None,
                                   '2.400' if row == 2 else None)
            elif column == 1 and row == 4:
                normalized = Value('missing', unit='RUB', notes=('missing_not_zero',))
            else:
                normalized = Value('text', raw)
            cells.append(DataCell(address, row, column, raw, normalized, source))
    dataset = Dataset('synthetic-ledger', 'Контроль данных проекта', DatasetPolicy(), tuple(cells), (),
                      tuple(c.source for c in cells), limitations=('Учебные данные',))
    return dataset, values


def synthetic_spec():
    return dict(contract='artifact-1', title='Обзор данных проекта', format='comparison', elements=[
        dict(id='adjustment', label='Корректировка', status='observed',
             text='Точное отрицательное значение сохранено без округления.',
             quantity=dict(value='-0.000001', unit='RUB'),
             proof=dict(kind='dataset', dataset_id='synthetic-ledger', address='B2')),
        dict(id='measurement', label='Наблюдение с диапазоном', status='confirmed',
             quantity=dict(value='2.345', lower='2.300', upper='2.400', unit='RUB'),
             proof=dict(kind='dataset', dataset_id='synthetic-ledger', address='B3')),
        dict(id='missing', label='Ожидаемые данные', status='unknown',
             text='Значение отсутствует и не заменяется нулём.'),
        dict(id='review', label='Следующий шаг', status='proposed',
             text='Уточнить недостающий показатель перед сравнением итогов.'),
    ], relations=[dict(id='next', **{'from': 'missing', 'to': 'review'}, kind='dependency',
                       label='Предложенный шаг зависит от получения данных')],
        questions=['Когда появится исходное значение?'], style={})


class DocumentVisualTests(unittest.TestCase):
    def test_docx_hierarchy_palette_sources_and_exact_numbers(self):
        from docx import Document
        spec = synthetic_spec()
        before = deepcopy(spec)
        raw = report_docx(spec, 3, {})
        doc = Document(BytesIO(raw))
        self.assertEqual(before, spec)
        self.assertEqual('Title', doc.paragraphs[0].style.name)
        self.assertEqual('300000', str(doc.styles['Title'].font.color.rgb))
        self.assertEqual([], doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'))
        self.assertEqual('840018', str(doc.styles['Heading 1'].font.color.rgb))
        text = '\n'.join(p.text for p in doc.paragraphs)
        for expected in ('-0.000001 RUB', '2.345 RUB', '2.300 … 2.400 RUB', 'Значение не задано',
                         'Источники', 'Открытые вопросы', canonical(spec['elements'][0]['proof'])):
            self.assertIn(expected, text)
        self.assertIn('unknown', doc.tables[0].cell(3, 1).text)
        self.assertEqual(5, len(doc.tables[0].rows))
        with ZipFile(BytesIO(raw)) as archive:
            body = archive.read('word/document.xml').decode()
            footer = archive.read('word/footer1.xml').decode()
        self.assertIn('w:tblHeader', body)
        self.assertIn('C0A8A8', body)
        self.assertIn('E4D8D8', body)
        self.assertIn('PAGE', footer)
        self.assertIn('NUMPAGES', footer)
        self.assertNotIn('E47854', body)
        self.assertNotIn('403020', body)

    def test_workbook_exact_literals_provenance_and_visual_contract(self):
        from openpyxl import load_workbook
        data, matrix = synthetic_dataset()
        before = data.to_dict()
        book = load_workbook(BytesIO(dataset_xlsx(data, matrix)))
        main, sources = book['Verified data'], book['Sources']
        self.assertEqual(before, data.to_dict())
        self.assertEqual('-0.000001', main['B2'].value)
        self.assertEqual('0', main['B4'].value)
        self.assertEqual('нет данных', main['B5'].value)
        self.assertEqual('2.345', main['B3'].value)
        self.assertIn('"lower":"2.300"', main['B3'].comment.text)
        self.assertIn('missing_not_zero', main['B5'].comment.text)
        for row, item in enumerate(data.cells, 2):
            self.assertEqual(item.raw, sources.cell(row, 2).value)
            self.assertEqual(canonical(asdict(item.source)), sources.cell(row, 5).value)
            self.assertEqual(item.normalized.unit, sources.cell(row, 4).value)
        for sheet in book:
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual('A2', sheet.freeze_panes)
            self.assertEqual('$1:$1', sheet.print_title_rows)
            self.assertIsNotNone(sheet.auto_filter.ref)
            self.assertEqual(1, sheet.page_setup.fitToWidth)
            self.assertEqual(0, sheet.page_setup.fitToHeight)
            self.assertEqual('00840018', sheet['A1'].fill.fgColor.rgb)
            self.assertEqual('00E4D8D8', sheet['A2'].fill.fgColor.rgb)
            self.assertGreater(sheet.column_dimensions['B'].width, 10)
            for row in sheet:
                for cell in row:
                    self.assertNotEqual('f', cell.data_type)
                    if cell.value is not None:
                        self.assertEqual('s', cell.data_type)
        self.assertEqual(data.name, main.oddHeader.center.text)

    def test_header_row_policy_preserves_source_positions(self):
        from openpyxl import load_workbook
        data, matrix = synthetic_dataset()
        book = load_workbook(BytesIO(dataset_xlsx(replace(data, policy=DatasetPolicy(header_row=None)), matrix)))
        main = book['Verified data']
        self.assertIsNone(main.freeze_panes)
        self.assertIsNone(main.auto_filter.ref)
        self.assertNotEqual('00840018', main['A1'].fill.fgColor.rgb)
        book = load_workbook(BytesIO(dataset_xlsx(replace(data, policy=DatasetPolicy(header_row=2)), matrix)))
        self.assertEqual('A4', book['Verified data'].freeze_panes)
        self.assertEqual('$3:$3', book['Verified data'].print_title_rows)
        self.assertEqual('-0.000001', book['Verified data']['B2'].value)

    def test_wide_workbook_prints_legibly_on_horizontal_pages(self):
        from artifacts.document_styles import _style_sheet
        from artifacts.styles import semantic_theme
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['Длинный заголовок столбца ' + str(index) for index in range(20)])
        sheet.append(['точное значение'] * 20)
        _style_sheet(sheet, semantic_theme(), 'Широкие данные', 1)
        self.assertEqual('landscape', sheet.page_setup.orientation)
        self.assertFalse(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(0, sheet.page_setup.fitToWidth)
        self.assertEqual(100, sheet.page_setup.scale)
        self.assertEqual('$A:$A', sheet.print_title_cols)

    def test_title_header_commands_and_native_formula_remain_literal(self):
        from openpyxl import load_workbook
        data, matrix = synthetic_dataset()
        cells = list(data.cells)
        cells[4] = replace(cells[4], raw='=1+1', normalized=Value('formula'), formula='=1+1', cached_raw='2')
        matrix[1][1] = '=1+1'
        raw = dataset_xlsx(replace(data, name='Название &P', cells=tuple(cells)), matrix)
        book = load_workbook(BytesIO(raw))
        self.assertEqual('Название &&P', book['Verified data'].oddHeader.center.text)
        self.assertEqual('=1+1', book['Verified data']['B2'].value)
        self.assertEqual('s', book['Verified data']['B2'].data_type)
        self.assertIn('не выполняется', book['Verified data']['B2'].comment.text)
        self.assertEqual('=1+1', book['Sources']['B6'].value)
        self.assertEqual('s', book['Sources']['B6'].data_type)

    def test_oversized_cells_fail_instead_of_silent_truncation(self):
        from openpyxl import load_workbook
        data, matrix = synthetic_dataset()
        cells = list(data.cells)
        matrix[1][2] = 'я' * 32767
        cells[5] = replace(cells[5], raw=matrix[1][2], normalized=Value('text', matrix[1][2]))
        book = load_workbook(BytesIO(dataset_xlsx(replace(data, cells=tuple(cells)), matrix)))
        self.assertEqual(matrix[1][2], book['Verified data']['C2'].value)
        matrix[1][2] += 'я'
        with self.assertRaisesRegex(MaterialError, 'document_cell_text_budget'):
            dataset_xlsx(data, matrix)
        # A normalized number is short, while its raw source or source locator
        # can still exceed the budget on the Sources worksheet.
        data, matrix = synthetic_dataset()
        cells = list(data.cells)
        cells[4] = replace(cells[4], raw='x' * 32768)
        with self.assertRaisesRegex(MaterialError, 'document_cell_text_budget'):
            dataset_xlsx(replace(data, cells=tuple(cells)), matrix)
        cells[4] = replace(data.cells[4], source=replace(data.cells[4].source, asset_id='x' * 32768))
        with self.assertRaisesRegex(MaterialError, 'document_cell_text_budget'):
            dataset_xlsx(replace(data, cells=tuple(cells)), matrix)

    def test_docx_embedded_visuals_are_fitted_and_decorations_labeled(self):
        from docx import Document
        from PIL import Image
        out = BytesIO()
        Image.new('RGB', (200, 1200), '#E4D8D8').save(out, 'PNG')
        doc = Document(BytesIO(report_docx(synthetic_spec(), 1, {'illustration-1.png': out.getvalue()})))
        self.assertLessEqual(doc.inline_shapes[0].height / 914400, 7.851)
        self.assertIn('Художественная иллюстрация; не свидетельство фактов.', '\n'.join(p.text for p in doc.paragraphs))


class DocumentToolVisualTests(unittest.IsolatedAsyncioTestCase):
    async def test_xlsx_tool_keeps_csv_bytes_and_semantic_hash(self):
        import csv
        from io import StringIO
        from agents.tools.documents import register_documents
        from agents.tools.registry import Registry, ToolContext
        from agents.subscriptions import semantic_fingerprint
        data, matrix = synthetic_dataset()
        registry = Registry()
        register_documents(registry)
        context = ToolContext(None, SimpleNamespace(repository=SimpleNamespace(pool=None)), 'p', 't', 'k', AsyncMock())
        with patch('materials.dataset_repository.DatasetRepository.load_dataset', AsyncMock(return_value=data)):
            csv_result = await registry.call('documents.data_export', dict(dataset_id=data.id, format='csv'), context)
            xlsx_result = await registry.call('documents.data_export', dict(dataset_id=data.id, format='xlsx'), context)
        expected = StringIO(newline='')
        csv.writer(expected).writerows([
            ["'" + str(v) if isinstance(v, str) and v[:1] in ('=', '+', '-', '@') and
             not any(x.row == r and x.column == col and x.normalized.kind == 'number' for x in data.cells)
             else v for col, v in enumerate(row)] for r, row in enumerate(matrix)])
        self.assertEqual(expected.getvalue().encode('utf-8-sig'), base64.b64decode(csv_result.outputs['files'][0]['base64']))
        self.assertEqual(semantic_fingerprint(matrix), csv_result.outputs['content_hash'])
        self.assertEqual(csv_result.outputs['content_hash'], xlsx_result.outputs['content_hash'])
        self.assertEqual((data.id,), xlsx_result.dependencies)


def write_samples(destination):
    from artifacts.export import export
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    spec = ArtifactSpec(synthetic_spec())
    data, matrix = synthetic_dataset()
    (destination / 'synthetic-report.docx').write_bytes(report_docx(spec.to_dict(), 1, export(spec)))
    (destination / 'synthetic-data.xlsx').write_bytes(dataset_xlsx(data, matrix))
    return destination


if __name__ == '__main__':
    import sys
    if len(sys.argv) == 3 and sys.argv[1] == '--samples':
        print(write_samples(sys.argv[2]))
    else:
        unittest.main()
