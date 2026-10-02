import base64,asyncio,csv
from io import BytesIO,StringIO
from dataclasses import asdict
from hashlib import sha256
from agents.tools.registry import Tool,ToolResult,object_schema,STRING
from materials.types import MaterialError
from agents.subscriptions import semantic_fingerprint

def register_documents(registry):
    async def data_export(args,c):
        from materials.dataset_repository import DatasetRepository
        data=await DatasetRepository(c.service.repository).load_dataset(args['dataset_id'],c.actor)
        cells=data.cells; maxr=max(x.row for x in cells)+1; maxc=max(x.column for x in cells)+1
        matrix=[['' for _ in range(maxc)] for _ in range(maxr)]
        for x in cells: matrix[x.row][x.column]=x.normalized.value if x.normalized.kind=='number' else x.raw
        if args['format']=='csv':
            # Spreadsheet formula injection is escaped only in exported presentation.
            out=StringIO(newline=''); writer=csv.writer(out)
            writer.writerows(["'"+str(v) if isinstance(v,str) and v[:1] in ('=','+','-','@') and not any(x.row==r and x.column==j and x.normalized.kind=='number' for x in cells) else v for j,v in enumerate(row)] for r,row in enumerate(matrix))
            content=out.getvalue().encode('utf-8-sig')
        else:
            from openpyxl import Workbook
            wb=Workbook(); ws=wb.active; ws.title='Verified data'
            for r,row in enumerate(matrix,1):
                for col,v in enumerate(row,1):
                    cell=ws.cell(r,col,str(v) if v is not None else ''); cell.data_type='s'
            provenance=wb.create_sheet('Sources'); provenance.append(['Cell','Original','Normalized','Unit','Evidence'])
            from materials.types import canonical
            for x in cells: provenance.append([x.address,str(x.raw),str(x.normalized.value),x.normalized.unit,canonical(asdict(x.source))])
            out=BytesIO(); wb.save(out); content=out.getvalue()
        return ToolResult('success',dict(files=[dict(name='data.'+args['format'],base64=base64.b64encode(content).decode(),sha256=sha256(content).hexdigest())],content_hash=semantic_fingerprint(matrix)),dependencies=(data.id,))
    registry.register(Tool('documents.data_export','1',object_schema(dict(dataset_id=STRING,format=dict(enum=['csv','xlsx']))),object_schema(dict(files=dict(type='array'),content_hash=STRING)),data_export,max_bytes=4000000))
    async def report(args,c):
        from artifacts.revisions import ArtifactRepository
        from artifacts.spec import ArtifactSpec
        row=await ArtifactRepository(c.service.repository).get(args['artifact_id'],c.actor)
        if row['revision']!=args['revision']: raise MaterialError('stale_artifact_revision')
        if args['format']=='pdf':
            from artifacts.export import export_current
            repo=ArtifactRepository(c.service.repository); repo.service=c.service
            content=(await export_current(repo,args['artifact_id'],c.actor,revision=args['revision']))[0]['report.pdf']
        else:
            from docx import Document
            from docx.shared import Inches
            from materials.types import canonical
            doc=Document(); doc.add_heading(row['spec']['title'],0)
            table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
            for cell,text in zip(table.rows[0].cells,['Элемент','Статус','Значение']): cell.text=text
            for e in row['spec']['elements']:
                cells=table.add_row().cells; cells[0].text=e['label']; cells[1].text=e['status']; cells[2].text=(e['quantity']['value']+' '+e['quantity']['unit']) if e.get('quantity') else 'Значение не задано'
                doc.add_heading(e['label'],1); doc.add_paragraph(e['status']+' / '+e.get('text',''))
                if e.get('quantity'): doc.add_paragraph(canonical(e['quantity']))
                if e.get('proof'): doc.add_paragraph('Source: '+canonical(e['proof']))
            for r in row['spec'].get('relations',[]): doc.add_paragraph(canonical(r))
            from artifacts.export import export_current
            repo=ArtifactRepository(c.service.repository); repo.service=c.service
            files,_=await export_current(repo,args['artifact_id'],c.actor,revision=args['revision'])
            doc.add_heading('Визуальный результат',1)
            for name,pixels in files.items():
                if name.endswith('.png'):
                    doc.add_picture(BytesIO(pixels),width=Inches(6))
                    doc.add_paragraph('Художественная иллюстрация; не свидетельство фактов.' if name.startswith('illustration-') else 'Программный рендер проверяемой спецификации.')
            out=BytesIO(); doc.save(out); content=out.getvalue()
        return ToolResult('success',dict(files=[dict(name='report.'+args['format'],base64=base64.b64encode(content).decode(),sha256=sha256(content).hexdigest())],content_hash=semantic_fingerprint(row['spec'])),dependencies=(row['head'],))
    registry.register(Tool('documents.report','1',object_schema(dict(artifact_id=STRING,revision=dict(type='integer',minimum=1),format=dict(enum=['pdf','docx']))),object_schema(dict(files=dict(type='array'),content_hash=STRING)),report,timeout=120,max_bytes=4000000))
