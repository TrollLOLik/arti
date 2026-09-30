"""Bounded CSV/XLSX extraction. Formulas and caches are observations, not results."""
import csv
from datetime import datetime,date
from hashlib import sha256
from io import BytesIO,StringIO
import importlib.metadata
import re
import time
from xml.etree import ElementTree as ET
import zipfile
from materials.extractors.isolation import WorkerLimits,run_worker
from materials.types import ContentBlock,ExtractionBundle,ExtractionManifest,Locator,MaterialError,block_id,canonical
from materials.validation import inspect_bytes

XLSX='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/package/2006/relationships'}


class TableExtractor:
    version='spreadsheet-source-1'
    def __init__(self,max_cells=3000,max_columns=128,max_sheets=16,delimiter=None,encoding='utf-8-sig'):
        if not 1<=max_cells<=3800 or not 1<=max_columns<=256 or not 1<=max_sheets<=32 or delimiter not in (None,',',';','\t','|') or encoding not in ('utf-8-sig','cp1251'):
            raise MaterialError('invalid_table_options')
        self.options=dict(max_cells=max_cells,max_columns=max_columns,max_sheets=max_sheets,delimiter=delimiter,encoding=encoding)
        self.limits=WorkerLimits()
    @property
    def cache_version(self):
        return self.version+':'+sha256(canonical([self.options,importlib.metadata.version('openpyxl')]).encode()).hexdigest()[:24]
    async def extract_async(self,aid,version,data,mime):
        result=await run_worker(dict(operation='table_extract',options=self.options,asset_id=aid,version=version,mime=mime),data,self.limits)
        return ExtractionBundle.from_dict(result)
    def extract(self,aid,version,data,mime):
        if self.options['encoding']=='cp1251' and mime=='text/csv':
            # Explicit encoding is accepted only in the isolated worker. Byte/MIME
            # intake remains UTF-8 unless importer supplies decoded UTF-8 data.
            if len(data)>10*1024**2: raise MaterialError('file_size_limit')
        else: inspect_bytes(data,'input.csv' if mime=='text/csv' else 'input.xlsx',mime)
        self.aid,self.revision=aid,version; self.blocks=[]; self.used=0; self.issues=[]
        self.deadline=time.monotonic()+70
        if mime=='text/csv': total,processed=self._csv(data); unit='row'
        elif mime==XLSX: total,processed=self._xlsx(data); unit='cell'
        else: raise MaterialError('table_mime_unsupported')
        coverage='complete' if total==processed and not any(s.endswith('_omitted') for s in self.issues) else 'partial'
        return ExtractionBundle(aid,version,self.cache_version,tuple(self.blocks),ExtractionManifest(total,processed,coverage,tuple(dict.fromkeys(self.issues)),unit))
    def _chunk(self,name,cells,metadata):
        if not cells: return
        # Chunk boundaries do not change cell addresses; dataset_group joins chunks.
        from openpyxl.utils import get_column_letter
        start,end=min(c['row'] for c in cells),max(c['row'] for c in cells)
        cols=sorted({c['column'] for c in cells}); last=max(cols)
        values={(c['row'],c['column']):c['text'] for c in cells}
        rows=[['' if (r,c) not in values else values[(r,c)] for c in range(last+1)] for r in range(start,end+1)]
        locator=Locator('cell',sheet=name,cell=f'A{start+1}:{get_column_letter(last+1)}{end+1}')
        ordinal=len(self.blocks); bid=block_id(self.aid,self.revision,'table',locator,ordinal)
        original=[]; children=[]
        for cell in cells:
            source=Locator('cell',sheet=name,cell=cell['address']); index=ordinal+1+len(children)
            cid=block_id(self.aid,self.revision,'text',source,index)
            original.append({**cell,'block_id':cid})
            children.append(ContentBlock(cid,'text',source,cell['text'],bid,index,quality='unassessed',metadata={**cell,'role':'table_cell','method':metadata['method']}))
        parent=ContentBlock(bid,'table',locator,'\n'.join('\t'.join(row) for row in rows),ordinal=ordinal,
            metadata={**metadata,'role':'table','dataset_group':name,'rows':rows,'cells':original})
        self.blocks.append(parent); self.blocks.extend(children)
    def _csv(self,data):
        try: text=data.decode(self.options['encoding'])
        except UnicodeError as exc: raise MaterialError('csv_encoding_required') from exc
        delimiter=self.options['delimiter']
        if delimiter is None:
            try: delimiter=csv.Sniffer().sniff(text[:16384],delimiters=',;\t|').delimiter
            except csv.Error: delimiter=','; self.issues.append('csv_delimiter_default_comma')
        rows=csv.reader(StringIO(text,newline=''),delimiter=delimiter,strict=True)
        total=0; processed=0; chunk=[]; first_width=None; stopped=False
        from openpyxl.utils import get_column_letter
        try:
            for ri,row in enumerate(rows):
                total+=1
                if stopped: continue  # Count remaining physical records with bounded memory.
                if first_width is None: first_width=len(row)
                width=max(first_width,len(row))
                if self.used+width>self.options['max_cells'] or width>self.options['max_columns'] or time.monotonic()>self.deadline:
                    self.issues.append('table_cell_budget_reached'); stopped=True; continue
                if len(row)!=first_width: self.issues.append('csv_ragged_rows')
                for ci in range(width):
                    raw=row[ci] if ci<len(row) else ''
                    if len(raw)>16000: raise MaterialError('cell_text_budget')
                    chunk.append(dict(row=ri,column=ci,address=get_column_letter(ci+1)+str(ri+1),raw=raw,text=raw,native_kind='string',
                        csv_record=ri+1,absent_in_record=ci>=len(row),formula_like_text=raw.startswith(('=','+','-','@'))))
                self.used+=width; processed+=1
                if len(chunk)>=200 or sum(len(canonical(c)) for c in chunk)>80000:
                    self._chunk('CSV',chunk,dict(method='csv_records',delimiter=delimiter,encoding=self.options['encoding'])); chunk=[]
        except csv.Error as exc: raise MaterialError('invalid_csv') from exc
        self._chunk('CSV',chunk,dict(method='csv_records',delimiter=delimiter,encoding=self.options['encoding']))
        self.issues.append('csv_formula_like_text_not_executed')
        return total,processed
    def _xlsx(self,data):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries,get_column_letter
        workbook=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        total=0; processed=0
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                names=archive.namelist()
                if any('vbaProject' in n for n in names): raise MaterialError('macros_not_supported')
                if any(n.startswith('xl/externalLinks/') for n in names): self.issues.append('external_links_not_followed')
                if any(n.startswith(('xl/drawings/','xl/charts/')) for n in names): self.issues.append('spreadsheet_visuals_omitted')
                # Use workbook relationships, not a guessed sheet1.xml path/order.
                book=ET.fromstring(archive.read('xl/workbook.xml'))
                relationships=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
                targets={r.attrib['Id']:r.attrib['Target'] for r in relationships}
                for si,sheet in enumerate(book.find('s:sheets',NS)):
                    name=sheet.attrib['name']; rid=sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
                    target=targets[rid]
                    path=target.lstrip('/') if target.startswith('/') else 'xl/'+target
                    if '..' in path.split('/'): raise MaterialError('invalid_workbook_relationship')
                    root=ET.fromstring(archive.read(path)); xml_cells=root.findall('s:sheetData/s:row/s:c',NS)
                    if si>=self.options['max_sheets']:
                        total+=len(xml_cells)
                        self.issues.append('sheet_budget_reached'); continue
                    ws=workbook[name]
                    merges=[n.attrib['ref'] for n in root.findall('s:mergeCells/s:mergeCell',NS)]
                    merge_bounds=[(r,range_boundaries(r)) for r in merges]
                    hidden_rows={int(n.attrib['r']) for n in root.findall('s:sheetData/s:row',NS) if n.attrib.get('hidden')=='1'}
                    hidden_cols=[(int(n.attrib['min']),int(n.attrib['max'])) for n in root.findall('s:cols/s:col',NS) if n.attrib.get('hidden')=='1']
                    # Iterate actual stored coordinates; declared worksheet dimensions
                    # can be stale or maliciously enormous. Read-only indexing remains
                    # costly, so consume rows once with bounded concrete coordinates.
                    coords={n.attrib['r']:n for n in xml_cells}
                    if not coords: continue
                    bounds=[range_boundaries(a) for a in coords]
                    max_row=max(b[1] for b in bounds); max_col=max(b[0] for b in bounds)
                    total+=max_row*max_col
                    if max_row*max_col>100000 or max_col>self.options['max_columns']:
                        self.issues.append('worksheet_sparse_dimension_budget'); continue
                    chunk=[]
                    for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=max_row,max_col=max_col)):
                        for ci,cell in enumerate(row):
                            if self.used>=self.options['max_cells'] or time.monotonic()>self.deadline:
                                self.issues.append('table_cell_budget_reached'); break
                            # EmptyCell carries no coordinate. Enumerated positions
                            # locate genuine blank workbook cells without fabricating a value.
                            address=get_column_letter(ci+1)+str(ri+1)
                            node=coords.get(address); value=node.find('s:v',NS) if node is not None else None
                            lexical=value.text if value is not None and value.text is not None else ''
                            kind={'n':'number','b':'boolean','e':'error','f':'formula','d':'date'}.get(cell.data_type,'string')
                            formula=cell.value if cell.data_type=='f' and isinstance(cell.value,str) else None
                            if cell.data_type=='f' and formula is None: kind='formula'; formula='[unsupported_array_formula]'
                            native_value=cell.value
                            if isinstance(native_value,(datetime,date)): kind='date'
                            raw=lexical if kind in ('number','boolean','error','formula') else str(native_value or '')
                            if formula: raw=formula
                            if len(raw)>16000: raise MaterialError('cell_text_budget')
                            merged=next((r for r,(x0,y0,x1,y1) in merge_bounds if x0<=ci+1<=x1 and y0<=ri+1<=y1),None)
                            if merged and range_boundaries(merged)[:2]!=(ci+1,ri+1): kind='merged'
                            iso=native_value.isoformat() if isinstance(native_value,(datetime,date)) else None
                            if kind=='date' and lexical=='60' and not workbook.epoch.year==1904:
                                kind='error'; self.issues.append('excel_1900_phantom_leap_day')
                            cache_raw=lexical if formula and value is not None else None
                            text=formula or raw
                            if formula: text += ' [cached='+str(cache_raw)+'; not recalculated]'
                            chunk.append(dict(row=ri,column=ci,address=address,raw=raw,text=text,native_kind=kind,absent_in_xml=node is None,
                                formula=formula,cached_raw=cache_raw,cache_status='unverified' if cache_raw is not None else 'missing',
                                source_lexical=lexical,iso_value=iso,number_format=cell.number_format,merged_range=merged,
                                hidden_row=ri+1 in hidden_rows,hidden_column=any(a<=ci+1<=b for a,b in hidden_cols)))
                            self.used+=1; processed+=1
                        if len(chunk)>=200 or sum(len(canonical(c)) for c in chunk)>80000:
                            self._chunk(name,chunk,dict(method='xlsx_cells',sheet_state=ws.sheet_state)); chunk=[]
                        if self.used>=self.options['max_cells'] or time.monotonic()>self.deadline: break
                    self._chunk(name,chunk,dict(method='xlsx_cells',sheet_state=ws.sheet_state))
                self.issues.append('xlsx_caches_not_verified')
        finally: workbook.close()
        return total,processed
