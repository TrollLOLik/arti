"""Structured PDF/DOCX + local OCR with original-coordinate evidence.

Layout, captions and table continuation are labelled heuristics. Native glyphs
and OCR words retain separate methods; omitted pages never count as processed.
"""
from dataclasses import replace
from io import BytesIO
import base64
import os
import re
import time
from materials.extractors.basic import BasicExtractor
from materials.extractors.isolation import WorkerLimits, run_worker
from materials.extractors.ocr import OCR, fingerprint, installation
from materials.types import ContentBlock, ExtractionBundle, ExtractionManifest, Locator, MaterialError, block_id

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def union(boxes):
    return [min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)]


def inside(inner, outer):
    x,y = (inner[0]+inner[2])/2,(inner[1]+inner[3])/2
    return outer[0] <= x <= outer[2] and outer[1] <= y <= outer[3]


def reading_order(items):
    """Whitespace XY cut. A spanning heading splits before a column cut."""
    if len(items)<2:
        return items
    def gap(axis):
        intervals = sorted((i['bbox'][axis],i['bbox'][axis+2]) for i in items)
        edge=intervals[0][1]; best=(0,None)
        for low,high in intervals[1:]:
            if low-edge>best[0]: best=(low-edge,(low+edge)/2)
            edge=max(edge,high)
        return best
    gx,cx = gap(0); gy,cy = gap(1)
    axis,cut = (0,cx) if gx>.035 else ((1,cy) if gy>.018 else (None,None))
    if axis is None:
        return sorted(items,key=lambda i:(round(i['bbox'][1],3),i['bbox'][0]))
    before=[i for i in items if i['bbox'][axis+2]<=cut]
    after=[i for i in items if i['bbox'][axis]>=cut]
    if len(before)+len(after)!=len(items) or not before or not after:
        return sorted(items,key=lambda i:(i['bbox'][1],i['bbox'][0]))
    return reading_order(before)+reading_order(after)


def lines(words):
    groups=[]
    for word in sorted(words,key=lambda w:(w['bbox'][1],w['bbox'][0])):
        centre=(word['bbox'][1]+word['bbox'][3])/2
        group=next((g for g in reversed(groups[-8:]) if abs(g['centre']-centre)<max(.004,(word['bbox'][3]-word['bbox'][1])*.5)),None)
        if group is None:
            group=dict(centre=centre,words=[]); groups.append(group)
        group['words'].append(word)
    result=[]
    for group in groups:
        chunks=[[]]
        for word in sorted(group['words'],key=lambda w:w['bbox'][0]):
            if chunks[-1] and word['bbox'][0]-chunks[-1][-1]['bbox'][2]>.045:
                chunks.append([])
            chunks[-1].append(word)
        for chunk in chunks:
            result.append(dict(text=' '.join(w['text'] for w in chunk),bbox=union([w['bbox'] for w in chunk]),words=chunk))
    return reading_order(result)


def ocr_lines(words):
    groups={}
    for word in words:
        groups.setdefault(tuple(word['group']),[]).append(word)
    result=[]
    for group in groups.values():
        ordered=sorted(group,key=lambda w:w['reading_bbox'][0])
        result.append(dict(text=' '.join(w['text'] for w in ordered),bbox=union([w['bbox'] for w in ordered]),
            reading_bbox=union([w['reading_bbox'] for w in ordered]),words=ordered))
    entries=reading_order([{**line,'bbox':line['reading_bbox'],'original_bbox':line['bbox']} for line in result])
    return [{**line,'bbox':line['original_bbox']} for line in entries]


class DocumentExtractor:
    version = 'structured-document-2'

    def __init__(self, max_units=64, max_chars=300_000, max_blocks=3500,
                 ocr_enabled=True, executable=None, tessdata=None, languages='rus+eng',
                 stack_fingerprint=None, wall_seconds=90, memory_mb=768):
        if not 1<=max_units<=256 or not 1<=max_chars<=2_000_000 or not 10<=max_blocks<=4000:
            raise MaterialError('invalid_extractor_budget')
        command,directory=installation()
        self.options=dict(max_units=max_units,max_chars=max_chars,max_blocks=max_blocks,ocr_enabled=ocr_enabled,
            executable=command if executable is None else executable,tessdata=directory if tessdata is None else tessdata,
            languages=languages,wall_seconds=wall_seconds,memory_mb=memory_mb)
        self.options['stack_fingerprint']=stack_fingerprint or fingerprint(self.options['executable'],self.options['tessdata'],languages)
        self.limits=WorkerLimits(wall_seconds=wall_seconds,memory_mb=memory_mb)
        self.ocr=OCR(self.options['executable'],self.options['tessdata'],languages,timeout=min(20,wall_seconds))

    @property
    def cache_version(self):
        from materials.types import canonical
        from hashlib import sha256
        return self.version+':'+sha256(canonical(self.options).encode()).hexdigest()[:24]

    async def extract_async(self, aid, version, data, mime):
        result=await run_worker(dict(operation='extract',options=self.options,asset_id=aid,version=version,mime=mime),data,self.limits)
        return ExtractionBundle.from_dict(result)

    def extract(self, aid, version, data, mime):
        # Synchronous entry is for the already isolated child and fixture tools.
        from materials.validation import inspect_bytes
        inspect_bytes(data,filename='original.txt' if mime.startswith('text/') else 'original.json',declared_mime=mime)
        if mime.startswith('text/') or mime=='application/json':
            return replace(BasicExtractor(self.options['max_units'],self.options['max_chars']).extract(aid,version,data,mime),extractor=self.cache_version)
        self.blocks=[]; self.issues=[]; self.used=0; self.aid=aid; self.asset_version=version
        self.deadline=time.monotonic()+self.limits.wall_seconds*.8
        self.ocr.deadline=self.deadline
        if mime=='application/pdf':
            total,processed=self._pdf(data); unit='page'
        elif mime==DOCX:
            total,processed=self._docx(data); unit='body_node'
        else:
            raise MaterialError('extractor_mime_unsupported')
        self._relations()
        partial=processed<total or any(i.startswith(('page_failed:','ocr_disabled','ocr_unavailable','ocr_failed','ocr_timeout','ocr_deadline','content_budget','docx_part_failed','docx_notes_omitted','docx_textboxes_omitted','page_visual_unread','nested_table_omitted','table_cell_budget')) for i in self.issues)
        return ExtractionBundle(aid,version,self.cache_version,tuple(self.blocks),
            ExtractionManifest(total,processed,'partial' if partial else 'complete',tuple(dict.fromkeys(self.issues)),unit))

    def _add(self,kind,locator,text='',*,parent=None,metadata=None,quality='unassessed',limitations=()):
        from materials.types import canonical
        metadata=metadata or {}
        if len(self.blocks)>=self.options['max_blocks'] or self.used+len(text)>self.options['max_chars'] or len(canonical(metadata))>190000:
            raise MaterialError('content_budget_reached')
        ordinal=len(self.blocks)
        block=ContentBlock(block_id(self.aid,self.asset_version,kind,locator,ordinal),kind,locator,text,parent,ordinal,
            'extracted',quality,tuple(limitations),metadata)
        self.blocks.append(block); self.used+=len(text)
        return block

    def _add_table(self,locator,text,*,parent=None,metadata=None,quality='unassessed'):
        table=self._add('table',locator,text,parent=parent,metadata=metadata,quality=quality)
        cells=[]
        for cell in metadata.get('cells',[]):
            source=Locator('region',page=locator.page,bbox=tuple(cell['bbox'])) if cell.get('bbox') else locator
            value=self._add('text',source,cell.get('text') or '',parent=table.block_id,
                metadata={**cell,'role':'table_cell','method':metadata['method'],'ocr':metadata.get('ocr',{})},quality=cell.get('quality',quality))
            cells.append({**cell,'block_id':value.block_id})
        updated=replace(table,metadata={**metadata,'cells':cells})
        self.blocks=[updated if b.block_id==table.block_id else b for b in self.blocks]
        return updated

    def _pdf(self,data):
        import pdfplumber
        import pypdfium2 as pdfium
        import pypdf
        self.issues.append('pdf_layout_heuristic_borderless_tables_unassessed')
        damaged=b'%%EOF' not in data[-2048:]
        parse_data=data+b'\n%%EOF\n' if damaged and b'startxref' in data[-2048:] else data
        if damaged: self.issues.append('pdf_recovered_missing_eof')
        probe=pypdf.PdfReader(BytesIO(parse_data),strict=False)
        if probe.is_encrypted:
            raise MaterialError('encrypted_document')
        with pdfium.PdfDocument(parse_data) as renderer, pdfplumber.open(BytesIO(parse_data)) as document:
            total=len(renderer); processed=0
            if len(document.pages)!=total: self.issues.append('pdf_page_count_disagreement')
            for index in range(min(total,self.options['max_units'])):
                if time.monotonic()>=self.deadline:
                    self.issues.append('parser_soft_deadline_reached'); break
                page_num=index+1
                try:
                    root=self._add('page',Locator('page',page=page_num),metadata=dict(role='page',coordinate_space='displayed_original_normalized'))
                except MaterialError as exc:
                    self.issues.append(exc.code); break
                try:
                    page=document.pages[index]
                    complete=self._pdf_page(page,renderer[index],page_num,root.block_id)
                    if complete: processed+=1
                    page.close()
                except MaterialError as exc:
                    self.issues.append(exc.code+':'+str(page_num))
                    if exc.code=='content_budget_reached': break
                except Exception:
                    self.issues.append('page_failed:'+str(page_num))
            if total>self.options['max_units']: self.issues.append('page_budget_reached')
        return total,processed

    def _pdf_page(self,page,render,page_num,parent):
        width,height=page.width,page.height
        def norm(b):
            return [max(0,b[0]/width),max(0,b[1]/height),min(1,b[2]/width),min(1,b[3]/height)]
        def loc(box): return Locator('region',page=page_num,bbox=tuple(box))
        words=[dict(text=w['text'],bbox=norm((w['x0'],w['top'],w['x1'],w['bottom'])),size=w.get('size',10),font=w.get('fontname',''))
               for w in page.dedupe_chars().extract_words(extra_attrs=['size','fontname'])]
        native_tables=[]
        for table in page.find_tables():
            rows=table.extract(); cells=[]
            if sum(len(row) for row in rows)>500: self.issues.append('table_cell_budget:'+str(page_num)); continue
            for ri,row in enumerate(table.rows):
                for ci,box in enumerate(row.cells):
                    if box:
                        cells.append(dict(row=ri,column=ci,text=rows[ri][ci] or '',bbox=norm(box)))
            native_tables.append(norm(table.bbox))
            self._add_table(loc(norm(table.bbox)),'\n'.join('\t'.join(c or '' for c in row) for row in rows),parent=parent,
                metadata=dict(role='table',rows=rows,cells=cells,method='native_grid',structure_quality='heuristic'))
        native_lines=lines([w for w in words if not any(inside(w['bbox'],t) for t in native_tables)])
        sizes=sorted(w['size'] for w in words); typical=sizes[len(sizes)//2] if sizes else 10
        for line in native_lines:
            text=line['text']; role='body'
            if re.match(r'^(рис(?:унок)?[. ]|fig(?:ure)?[. ]|таблица\s)',text,re.I): role='caption'
            elif re.match(r'^(?:[•●▪–-]|\d+[.)])\s',text): role='list_item'
            elif max(w['size'] for w in line['words'])>typical*1.18: role='heading'
            self._add('text',loc(line['bbox']),text,parent=parent,
                metadata=dict(role=role,method='native_glyphs',words=[dict(text=w['text'],bbox=w['bbox']) for w in line['words']],layout_quality='heuristic'))
        image_boxes=[]
        for index,image in enumerate(page.images):
            box=norm((image['x0'],image['top'],image['x1'],image['bottom']))
            if box[0]>=box[2] or box[1]>=box[3]: continue
            image_boxes.append(box)
            self._add('image',loc(box),parent=parent,metadata=dict(role='embedded_image',method='pdf_image_region',image_index=index,
                resource=dict(kind='pdf_region',page=page_num,bbox=box),visual_content='uninterpreted'))
        # OCR entire rendered page for scan/mixed image areas, then keep only new regions.
        needs_ocr=not words or bool(image_boxes)
        complete=True
        if needs_ocr:
            if not self.options['ocr_enabled']:
                self.issues.append('ocr_disabled:'+str(page_num)); return False
            try:
                bitmap=render.render(scale=min(3, (16_000_000/(width*height))**.5))
                try: image=bitmap.to_pil().copy()
                finally: bitmap.close()
                result=self.ocr.read(image)
                # A native text layer can be wrong (e.g. a stale hidden OCR
                # layer). Keep disagreements as uncertain alternatives, never
                # silently replace either source with a more fluent reading.
                conflicts=[]
                def glyph_key(text): return re.sub(r'\W','',text).casefold()
                for observed in result['words']:
                    match=next((w for w in words if inside(observed['bbox'],w['bbox'])),None)
                    if match and glyph_key(match['text'])!=glyph_key(observed['text']) and observed['confidence']>=70:
                        conflicts.append(dict(native_text=match['text'],ocr_text=observed['text'],bbox=observed['bbox'],engine_score=observed['confidence']))
                if conflicts:
                    self.issues.append('native_ocr_disagreement:'+str(page_num))
                    for conflict in conflicts[:100]:
                        self._add('text',loc(conflict['bbox']),'[Native/OCR disagreement] '+conflict['native_text']+' / '+conflict['ocr_text'],
                            parent=parent,metadata={**conflict,'role':'disagreement','method':'native_ocr_comparison'},quality='uncertain')
                # Preserve native text, don't copy OCR over it or infer absence from it.
                selected=[w for w in result['words'] if not any(inside(w['bbox'],l['bbox']) for l in native_lines)
                    and not any(inside(w['bbox'],t) for t in native_tables)
                    and (not words or any(inside(w['bbox'],b) for b in image_boxes))]
                ocr_tables=[]
                for table in result['tables']:
                    if any(inside(table['bbox'],t) for t in native_tables): continue
                    # Avoid claiming raster-only cells when native glyphs already cover them.
                    if any(inside(w['bbox'],table['bbox']) for w in words): continue
                    ocr_tables.append(table['bbox'])
                    self._add_table(loc(table['bbox']),'\n'.join('\t'.join(row) for row in table['rows']),parent=parent,
                        metadata={**table,'role':'table','ocr':{k:v for k,v in result.items() if k not in ('words','tables')}},quality='uncertain')
                for line in ocr_lines([w for w in selected if not any(inside(w['bbox'],t) for t in ocr_tables)]):
                    uncertain=[w for w in line['words'] if w['confidence']<70 or '\ufffd' in w['text']
                        or (re.fullmatch(r'[0-9OОоЗзI|]{2,}',w['text']) and re.search(r'[OОоЗзI|]',w['text']))]
                    role='caption' if re.match(r'^(рис(?:унок)?[. ]|fig(?:ure)?[. ]|таблица\s)',line['text'],re.I) else 'body'
                    self._add('text',loc(line['bbox']),line['text'],parent=parent,quality='uncertain' if uncertain else 'unassessed',
                        metadata=dict(role=role,method='tesseract',reading_bbox=line['reading_bbox'],words=[{k:v for k,v in w.items() if k not in ('box','original_bbox')} for w in line['words']],
                            ocr={k:v for k,v in result.items() if k not in ('words','tables')},
                            reread_regions=[w['bbox'] for w in uncertain]),limitations=('ocr_not_human_verified',))
                if not selected and not native_tables and not words:
                    self.issues.append('page_visual_unread:'+str(page_num)); complete=False
                self.issues.append('ocr_printed_text_only:'+str(page_num))
                image.close()
            except MaterialError as exc:
                self.issues.append(exc.code+':'+str(page_num)); complete=False
        return complete

    def _docx(self,data):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        document=Document(BytesIO(data))
        nodes=[n for n in document.element.body if n.tag in (qn('w:p'),qn('w:tbl'))]
        total=len(nodes); processed=0
        def paragraph(node,owner,number,role='body',path='body'):
            para=Paragraph(node,owner); text=para.text
            style=para.style.name if para.style else ''
            if style.lower().startswith('heading'): role='heading'
            elif 'caption' in style.lower() or re.match(r'^(рис(?:унок)?[. ]|таблица\s)',text,re.I): role='caption'
            elif node.find('.//'+qn('w:numPr')) is not None or 'list' in style.lower(): role='list_item'
            block=self._add('text',Locator('paragraph',paragraph=number),text,
                metadata=dict(role=role,style=style,xml_path=path,method='docx_xml',pagination='unknown'))
            for drawing_index,blip in enumerate(node.iter(qn('a:blip'))):
                relationship=blip.get(qn('r:embed'))
                part=owner.part.related_parts.get(relationship)
                if part is None:
                    self.issues.append('docx_part_failed:image'); continue
                self._add('image',Locator('paragraph',paragraph=number),parent=block.block_id,
                    metadata=dict(role='embedded_image',method='docx_relationship',xml_path=path,
                        resource=dict(kind='docx_image',part=str(owner.part.partname),relationship=relationship,drawing_index=drawing_index),
                        mime=part.content_type,visual_content='uninterpreted'))
        def table(node,number,path,parent=None,owner=document):
            rows=[]; cells=[]
            for ri,tr in enumerate(node.findall(qn('w:tr'))):
                row=[]; col=0
                before=tr.find('./'+qn('w:trPr')+'/'+qn('w:gridBefore'))
                if before is not None: col=int(before.get(qn('w:val'),0))
                for tc in tr.findall(qn('w:tc')):
                    span=tc.find('./'+qn('w:tcPr')+'/'+qn('w:gridSpan'))
                    colspan=int(span.get(qn('w:val'),1)) if span is not None else 1
                    merge=tc.find('./'+qn('w:tcPr')+'/'+qn('w:vMerge'))
                    continuation=merge is not None and merge.get(qn('w:val'),'continue')=='continue'
                    text='\n'.join(Paragraph(p,owner).text for p in tc.findall(qn('w:p')))
                    cells.append(dict(row=ri,column=col,colspan=colspan,text=text,merge='continue' if continuation else ('start' if merge is not None else 'none'),
                        xml_path=f'{path}/row:{ri}/cell:{col}'))
                    row.append('' if continuation else text); col+=colspan
                rows.append(row)
            block=self._add_table(Locator('paragraph',paragraph=number),'\n'.join('\t'.join(row) for row in rows),parent=parent,
                metadata=dict(role='table',rows=rows,cells=cells,method='docx_xml',xml_path=path,pagination='unknown'))
            for ri,tr in enumerate(node.findall(qn('w:tr'))):
                for ci,tc in enumerate(tr.findall(qn('w:tc'))):
                    for ni,nested in enumerate(tc.findall(qn('w:tbl'))): table(nested,number,f'{path}/row:{ri}/cell:{ci}/nested:{ni}',block.block_id,owner)
                    for pi,p in enumerate(tc.findall(qn('w:p'))):
                        if p.find('.//'+qn('a:blip')) is not None: paragraph(p,owner,number,path=f'{path}/row:{ri}/cell:{ci}/p:{pi}')
        for i,node in enumerate(nodes,1):
            if time.monotonic()>=self.deadline:
                self.issues.append('parser_soft_deadline_reached'); break
            if i>self.options['max_units']:
                self.issues.append('body_node_budget_reached'); break
            try:
                if node.tag==qn('w:p'): paragraph(node,document,i,path=f'body/{i}')
                else: table(node,i,f'body/{i}')
                processed+=1
            except MaterialError as exc:
                self.issues.append(exc.code); break
            except Exception:
                self.issues.append('docx_part_failed:'+str(i))
        # Distinct sections can share linked parts; read each physical part once.
        seen=set()
        for si,section in enumerate(document.sections):
            for role in ('header','footer'):
                for prefix in ('','first_page_','even_page_'):
                    part=getattr(section,prefix+role)
                    if not part._has_definition or str(part.part.partname) in seen: continue
                    seen.add(str(part.part.partname))
                    for pi,p in enumerate(part.paragraphs):
                        try: paragraph(p._p,part,total+si+1,role,f'{part.part.partname}/p:{pi}')
                        except MaterialError: self.issues.append('content_budget_reached'); break
                    for ti,t in enumerate(part.tables):
                        try: table(t._tbl,total+si+1,f'{part.part.partname}/tbl:{ti}',owner=part)
                        except MaterialError: self.issues.append('content_budget_reached'); break
        # Preserve honesty for OOXML content not represented by body nodes.
        import zipfile
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if 'word/footnotes.xml' in archive.namelist() or 'word/endnotes.xml' in archive.namelist():
                self.issues.append('docx_notes_omitted')
        if document.element.body.findall('.//'+qn('w:txbxContent')):
            self.issues.append('docx_textboxes_omitted')
        self.issues.append('docx_pagination_unknown')
        return total,processed

    def _relations(self):
        # Repeat-margin detection uses >=2 different pages, never body repetition.
        candidates={}
        for b in self.blocks:
            if b.kind=='text' and b.locator.page and b.locator.bbox and (b.locator.bbox[3]<.09 or b.locator.bbox[1]>.91):
                key=re.sub(r'\d+','#',b.text.strip())
                candidates.setdefault(key,set()).add(b.locator.page)
        replacements={}
        for b in self.blocks:
            metadata=dict(b.metadata)
            if len(candidates.get(re.sub(r'\d+','#',b.text.strip()),()))>=2 and b.locator.bbox and (b.locator.bbox[3]<.09 or b.locator.bbox[1]>.91):
                metadata.update(role='header' if b.locator.bbox[3]<.09 else 'footer',role_method='repeated_margin_heuristic')
            if b.kind=='image':
                captions=[c for c in self.blocks if c.metadata.get('role')=='caption' and c.locator.page==b.locator.page
                    and ((c.locator.bbox and b.locator.bbox and abs(c.locator.bbox[1]-b.locator.bbox[3])<.09)
                         or (c.locator.paragraph and b.locator.paragraph and abs(c.locator.paragraph-b.locator.paragraph)<=1))]
                if captions:
                    caption=min(captions,key=lambda c:abs((c.locator.bbox or [0,0,0,0])[1]-(b.locator.bbox or [0,0,0,0])[3]))
                    metadata.update(caption_id=caption.block_id,caption_relation='proximity_heuristic')
            replacements[b.block_id]=replace(b,metadata=metadata)
        self.blocks=[replacements[b.block_id] for b in self.blocks]
        # Conservatively link adjacent table fragments only with same header, shape
        # and page-boundary position; retain both originals and do not concatenate.
        tables=[b for b in self.blocks if b.kind=='table' and b.locator.page]
        updates={}
        for first in tables:
            for second in tables:
                if second.locator.page!=first.locator.page+1: continue
                rows1,rows2=first.metadata.get('rows',[]),second.metadata.get('rows',[])
                if rows1 and rows2 and rows1[0]==rows2[0] and first.locator.bbox[3]>.65 and second.locator.bbox[1]<.35:
                    updates[second.block_id]=replace(second,metadata={**second.metadata,'continues':first.block_id,'continuation_quality':'heuristic','repeated_header_rows':[0]})
        self.blocks=[updates.get(b.block_id,b) for b in self.blocks]
        # Root pages precede children; children follow inferred reading order.
        if any(b.kind=='page' for b in self.blocks):
            ordered=[]
            def append_tree(block):
                ordered.append(block)
                for child in [b for b in self.blocks if b.parent_id==block.block_id]:
                    append_tree(child)
            for root in [b for b in self.blocks if b.kind=='page']:
                children=[b for b in self.blocks if b.parent_id==root.block_id]
                entries=[dict(bbox=list(b.metadata.get('reading_bbox',b.locator.bbox)),block=b) for b in children if b.locator.bbox]
                ordered.append(root)
                for entry in reading_order(entries): append_tree(entry['block'])
            self.blocks=ordered

    async def region_async(self,data,mime,locator,*,reread=False,resource=None,orientation_hint=None):
        return await run_worker(dict(operation='region',options=self.options,mime=mime,locator=locator,reread=reread,resource=resource,orientation_hint=orientation_hint),data,self.limits)

    def region(self,data,mime,locator,reread=False,resource=None,orientation_hint=None):
        from PIL import Image
        from materials.validation import inspect_bytes
        inspect_bytes(data,filename='original',declared_mime=mime)
        if mime==DOCX and resource and resource.get('kind')=='docx_image':
            from docx import Document
            document=Document(BytesIO(data))
            owner=next((p for p in document.part.package.parts if str(p.partname)==resource['part']),None)
            part=owner.related_parts.get(resource['relationship']) if owner else None
            if part is None: raise MaterialError('docx_image_missing')
            inspect_bytes(part.blob,'image',part.content_type)
            image=Image.open(BytesIO(part.blob)); image.thumbnail((2400,2400))
            encoded,downsampled=self._encode_preview(image.convert('RGB')); image.close()
            return dict(mime='image/png',image_base64=encoded,locator=locator,preview_downsampled=downsampled)
        if mime!='application/pdf': raise MaterialError('region_mime_unsupported')
        import pypdfium2 as pdfium
        loc=Locator.from_dict(locator)
        if loc.kind.value!='region' or not loc.page: raise MaterialError('missing_page_region')
        with pdfium.PdfDocument(data) as document:
            if loc.page>len(document): raise MaterialError('invalid_page')
            page=document[loc.page-1]; width,height=page.get_size(); x0,y0,x1,y1=loc.bbox
            crop=(x0*width,(1-y1)*height,(1-x1)*width,y0*height)
            scale=min(4, (8_000_000/(width*height*(x1-x0)*(y1-y0)))**.5)
            bitmap=page.render(scale=scale,crop=crop)
            try: image=bitmap.to_pil().copy()
            finally: bitmap.close(); page.close()
        encoded,downsampled=self._encode_preview(image)
        result=dict(mime='image/png',image_base64=encoded,locator=locator,preview_downsampled=downsampled)
        if reread:
            observed=self.ocr.read(image,targeted=True,orientation_hint=orientation_hint)
            for word in observed['words']:
                box=word['bbox']; word['bbox']=[x0+box[0]*(x1-x0),y0+box[1]*(y1-y0),x0+box[2]*(x1-x0),y0+box[3]*(y1-y0)]
            result['observation']=observed
        image.close()
        return result

    @staticmethod
    def _encode_preview(image):
        # Preview budget is separate from OCR resolution. Keep original geometry.
        working=image.copy(); downsampled=False
        try:
            for _ in range(5):
                output=BytesIO(); working.save(output,format='PNG')
                if output.tell()<=5*1024**2:
                    return base64.b64encode(output.getvalue()).decode(),downsampled
                working.thumbnail((max(1,working.width//2),max(1,working.height//2)))
                downsampled=True
            raise MaterialError('preview_output_budget')
        finally:
            working.close()


def configured_extractor():
    if os.getenv('ARTI_DOCUMENTS_ENABLED','1').lower() in ('0','false','off'):
        return IsolatedNativeExtractor()
    return DocumentExtractor(ocr_enabled=os.getenv('ARTI_OCR_ENABLED','1').lower() not in ('0','false','off'))


class IsolatedNativeExtractor(BasicExtractor):
    """Rollback retains process limits even when structured layout/OCR is off."""
    version='isolated-native-document-1'

    async def extract_async(self,aid,version,data,mime):
        result=await run_worker(dict(operation='native_extract',options=dict(max_units=self.max_units,max_chars=self.max_chars),
            asset_id=aid,version=version,mime=mime),data,WorkerLimits())
        return ExtractionBundle.from_dict(result)
