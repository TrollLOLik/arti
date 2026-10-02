"""Immutable typed snapshots. Normalization never rewrites the original cell."""
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation,localcontext
from hashlib import sha256
import re
from materials.types import EvidenceRef, Locator, MaterialError, canonical

DATASET_VERSION='dataset-1'
UNITS={'1':('ratio','1'),'percent':('ratio','0.01'),'RUB':('currency:RUB','1'),
    'USD':('currency:USD','1'),'EUR':('currency:EUR','1'),'GBP':('currency:GBP','1'),
    'CAD':('currency:CAD','1'),'AUD':('currency:AUD','1'),'CNY':('currency:CNY','1'),'JPY':('currency:JPY','1'),
    'm':('length','1'),'cm':('length','0.01'),'km':('length','1000'),
    'g':('mass','0.001'),'kg':('mass','1'),'s':('time','1'),'min':('time','60'),'h':('time','3600'),
    'item':('count','1')}
ALIASES={'₽':'RUB','руб':'RUB','руб.':'RUB','рублей':'RUB','rub':'RUB',
    'usd':'USD','€':'EUR','eur':'EUR','gbp':'GBP','cad':'CAD','aud':'AUD','cny':'CNY','jpy':'JPY','%':'percent',
    'м':'m','см':'cm','км':'km','г':'g','кг':'kg','с':'s','сек':'s','мин':'min','ч':'h','шт':'item','шт.':'item'}


def unit_code(value):
    code=ALIASES.get(value.casefold(),value) if isinstance(value,str) else value
    if code not in UNITS: raise MaterialError('unknown_unit')
    return code


def decimal(value):
    try:
        number=Decimal(str(value))
        if not number.is_finite() or abs(number)>Decimal('1e100') or abs(number.as_tuple().exponent)>100 or len(number.as_tuple().digits)>100:
            raise MaterialError('numeric_budget')
        return number
    except (InvalidOperation,ValueError,TypeError) as exc:
        raise MaterialError('invalid_number') from exc


def decstr(value):
    number=decimal(value)
    if not number: return '0'
    return format(number,'f').rstrip('0').rstrip('.') if number.as_tuple().exponent<0 else format(number,'f')


def percent_format(number_format):
    # A quoted or escaped percent sign is a literal, not Excel's fraction display.
    return '%' in re.sub(r'"[^"]*"|\\.', '', number_format or '')


@dataclass(frozen=True)
class ColumnPolicy:
    index: int
    dtype: str='auto'
    unit: str='auto'
    locale: str='unknown'
    date_order: str='unknown'
    def __post_init__(self):
        if type(self.index) is not int or self.index<0 or self.dtype not in ('auto','number','text','date','boolean') or self.locale not in ('unknown','ru','en','de') or self.date_order not in ('unknown','DMY','MDY','YMD'):
            raise MaterialError('invalid_column_policy')
        if self.unit!='auto': unit_code(self.unit)


@dataclass(frozen=True)
class DatasetPolicy:
    header_row: int | None=0  # zero-based source row; explicit default assumption
    columns: tuple[ColumnPolicy,...]=()
    def __post_init__(self):
        if self.header_row is not None and (type(self.header_row) is not int or self.header_row<0): raise MaterialError('invalid_header_row')
        if len({c.index for c in self.columns})!=len(self.columns): raise MaterialError('duplicate_column_policy')
    def column(self,index): return next((c for c in self.columns if c.index==index),ColumnPolicy(index))
    @classmethod
    def from_dict(cls,value): return cls(value.get('header_row',0),tuple(ColumnPolicy(**c) for c in value.get('columns',())))


@dataclass(frozen=True)
class Value:
    kind: str
    value: str | None=None
    unit: str='1'
    lower: str | None=None
    upper: str | None=None
    notes: tuple[str,...]=()
    def __post_init__(self):
        if self.kind not in ('number','text','date','boolean','missing','ambiguous','invalid','formula','error','merged'):
            raise MaterialError('invalid_cell_kind')
        unit_code(self.unit)
        if self.kind=='number':
            value=decimal(self.value)
            low=decimal(self.lower if self.lower is not None else self.value)
            high=decimal(self.upper if self.upper is not None else self.value)
            if not low<=value<=high: raise MaterialError('invalid_value_interval')


def normalize(raw,policy=None,*,native_kind=None,iso_value=None):
    policy=policy or ColumnPolicy(0)
    text=('' if raw is None else str(raw)).strip(); unit='1' if policy.unit=='auto' else unit_code(policy.unit)
    if native_kind=='formula': return Value('formula',notes=('cached_value_not_evaluated',),unit=unit)
    if native_kind=='error': return Value('error',text,unit,notes=('source_spreadsheet_error',))
    if native_kind=='merged': return Value('merged',unit=unit,notes=('merged_continuation_not_zero',))
    if native_kind=='date' and iso_value:
        try: return Value('date',date.fromisoformat(iso_value[:10]).isoformat())
        except ValueError: return Value('invalid',notes=('invalid_native_date',))
    if native_kind=='boolean': return Value('boolean','true' if text in ('1','True','true') else 'false')
    if not text or text.casefold() in ('n/a','na','null','—','–','-','нет данных'):
        return Value('missing',unit=unit,notes=('missing_not_zero',))
    if policy.dtype=='text': return Value('text',text)
    if native_kind=='number':
        try: return Value('number',decstr(text),unit)
        except MaterialError: return Value('invalid',unit=unit,notes=('invalid_native_number',))
    if policy.dtype=='boolean':
        if text.casefold() in ('true','false','да','нет','1','0'):
            return Value('boolean','true' if text.casefold() in ('true','да','1') else 'false')
        return Value('invalid',notes=('invalid_boolean',))
    # ISO dates are unambiguous. Slash/dot dates require an order when both
    # interpretations are possible; invalid dates are never coerced to numbers.
    iso=re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})',text)
    short=re.fullmatch(r'(\d{1,2})[/.](\d{1,2})[/.](\d{4})',text)
    if iso or short:
        try:
            if iso: y,m,d=map(int,iso.groups())
            else:
                a,b,y=map(int,short.groups()); order=policy.date_order
                if order=='unknown':
                    if a>12: order='DMY'
                    elif b>12: order='MDY'
                    elif a==b: order='DMY'
                    else: return Value('ambiguous',notes=('date_order_required',))
                if order not in ('DMY','MDY'): return Value('invalid',notes=('invalid_date_order',))
                d,m=(a,b) if order=='DMY' else (b,a)
            return Value('date',date(y,m,d).isoformat())
        except ValueError: return Value('invalid',notes=('invalid_date',))
    if policy.dtype=='date': return Value('invalid',notes=('date_format_unrecognized',))
    # Explicit suffix/prefix units override a dimensionless column, but disagreeing
    # declarations remain invalid. Preserve differing currencies in the dataset.
    suffix=re.fullmatch(r'(.+?)\s*(₽|\$|€|£|%|RUB|USD|EUR|GBP|CAD|AUD|CNY|JPY|руб(?:\.|лей)?|кг|км|см|м|kg|km|cm|m|g|s|h|шт\.?)',text,re.I)
    prefix=re.fullmatch(r'(₽|\$|€|£)\s*(.+)',text)
    if suffix or prefix:
        numeric,mark=(suffix.group(1),suffix.group(2)) if suffix else (prefix.group(2),prefix.group(1))
        if mark in ('$','£'):
            allowed=('USD','CAD','AUD') if mark=='$' else ('GBP',)
            if unit=='1': return Value('ambiguous',notes=('currency_code_required',))
            if unit not in allowed: return Value('invalid',unit=unit,notes=('column_unit_conflict',))
            explicit=unit
        else: explicit=unit_code(mark)
        if unit!='1' and unit!=explicit: return Value('invalid',unit=unit,notes=('column_unit_conflict',))
        unit=explicit; text=numeric.strip()
    if '±' in text:
        parts=text.split('±')
        if len(parts)==2:
            center=normalize(parts[0],replace(policy,unit=unit)); margin=normalize(parts[1],replace(policy,unit=unit))
            if center.kind==margin.kind=='number' and decimal(margin.value)>=0:
                with localcontext() as context:
                    context.prec=100
                    n,e=decimal(center.value),decimal(margin.value)
                    return Value('number',decstr(n),unit,decstr(n-e),decstr(n+e),('explicit_interval_not_probability',))
        return Value('invalid',unit=unit,notes=('invalid_uncertainty_interval',))
    text=text.replace('\u2212','-').replace('\u00a0',' ').replace('\u202f',' ')
    negative=text.startswith('(') and text.endswith(')')
    if negative: text='-'+text[1:-1]
    if re.fullmatch(r'[+-]?0\d+',text) and policy.dtype=='auto': return Value('text',text,notes=('leading_zero_identifier',))
    if not re.fullmatch(r'[+-]?[0-9][0-9., ]*(?:[eE][+-]?\d+)?',text):
        return Value('invalid' if policy.dtype=='number' else 'text',text,unit,notes=('not_numeric',) if policy.dtype=='number' else ())
    try:
        numeric=text
        if ' ' in numeric:
            # Space grouping must use groups of three; arbitrary spaces aren't removed.
            if not re.fullmatch(r'[+-]?\d{1,3}(?: \d{3})+(?:[.,]\d+)?',numeric):
                return Value('invalid',unit=unit,notes=('invalid_digit_grouping',))
            numeric=numeric.replace(' ','')
        if policy.locale in ('en','de','ru'):
            decimal_mark=',' if policy.locale in ('de','ru') else '.'
            group='.' if policy.locale=='de' else (',' if policy.locale=='en' else None)
            if group and group in numeric:
                before=numeric.split(decimal_mark)[0]
                if not re.fullmatch(r'[+-]?\d{1,3}(?:'+re.escape(group)+r'\d{3})+',before):
                    return Value('invalid',unit=unit,notes=('locale_grouping_mismatch',))
                numeric=numeric.replace(group,'')
            if policy.locale=='ru' and '.' in numeric: return Value('invalid',unit=unit,notes=('locale_decimal_mismatch',))
            numeric=numeric.replace(decimal_mark,'.')
        elif ',' in numeric and '.' in numeric:
            mark=',' if numeric.rfind(',')>numeric.rfind('.') else '.'
            group='.' if mark==',' else ','
            before=numeric.split(mark)[0]
            if not re.fullmatch(r'[+-]?\d{1,3}(?:'+re.escape(group)+r'\d{3})+',before): return Value('ambiguous',unit=unit,notes=('mixed_separators',))
            numeric=numeric.replace(group,'').replace(mark,'.')
        elif ',' in numeric or '.' in numeric:
            mark=',' if ',' in numeric else '.'
            if re.fullmatch(r'[+-]?\d{1,3}'+re.escape(mark)+r'\d{3}',numeric):
                return Value('ambiguous',unit=unit,notes=('decimal_or_thousands',))
            if numeric.count(mark)>1:
                if not re.fullmatch(r'[+-]?\d{1,3}(?:'+re.escape(mark)+r'\d{3})+',numeric): return Value('invalid',unit=unit,notes=('invalid_digit_grouping',))
                numeric=numeric.replace(mark,'')
            else: numeric=numeric.replace(mark,'.')
        return Value('number',decstr(numeric),unit)
    except MaterialError:
        return Value('invalid',unit=unit,notes=('numeric_format_or_budget',))


@dataclass(frozen=True)
class DataCell:
    address: str
    row: int
    column: int
    raw: str
    normalized: Value
    source: EvidenceRef
    quality: str='unassessed'
    formula: str | None=None
    cached_raw: str | None=None
    metadata: dict=field(default_factory=dict)
    def __post_init__(self):
        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*',self.address) or self.row<0 or self.column<0 or self.quality not in ('verified','unassessed','uncertain','unreadable'):
            raise MaterialError('invalid_dataset_cell')
    @classmethod
    def from_dict(cls,value):
        source=value['source']; normalized=value['normalized']
        return cls(**{**value,'source':EvidenceRef(**{**source,'locator':Locator.from_dict(source['locator'])}),
            'normalized':Value(**{**normalized,'notes':tuple(normalized.get('notes',()))})})


@dataclass(frozen=True)
class Dataset:
    series_key: str
    name: str
    policy: DatasetPolicy
    cells: tuple[DataCell,...]
    columns: tuple[dict,...]
    source_refs: tuple[EvidenceRef,...]
    coverage: str='complete'
    limitations: tuple[str,...]=()
    previous_id: str | None=None
    corrections: tuple[dict,...]=()
    contract_version: str=DATASET_VERSION
    def __post_init__(self):
        if not self.cells or len(self.cells)>4000 or len({c.address for c in self.cells})!=len(self.cells) or not self.source_refs or self.contract_version!=DATASET_VERSION:
            raise MaterialError('invalid_dataset')
        if self.coverage not in ('complete','partial','unknown','failed'): raise MaterialError('invalid_dataset_coverage')
    @property
    def id(self): return sha256(canonical(self.to_dict()).encode()).hexdigest()
    def to_dict(self): return asdict(self)
    @classmethod
    def from_dict(cls,value):
        return cls(**{**value,'policy':DatasetPolicy.from_dict(value['policy']),
            'cells':tuple(DataCell.from_dict(c) for c in value['cells']),
            'source_refs':tuple(EvidenceRef(**{**r,'locator':Locator.from_dict(r['locator'])}) for r in value['source_refs']),
            'columns':tuple(value['columns']),'limitations':tuple(value.get('limitations',())),
            'corrections':tuple(value.get('corrections',()))})
    def cell(self,address):
        value=next((c for c in self.cells if c.address==address.upper()),None)
        if value is None: raise MaterialError('dataset_cell_missing')
        return value
    def select(self,selection):
        from openpyxl.utils.cell import range_boundaries
        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?',selection): raise MaterialError('invalid_selection')
        c0,r0,c1,r1=range_boundaries(selection)
        if c0>c1 or r0>r1 or (c1-c0+1)*(r1-r0+1)>4000: raise MaterialError('selection_budget')
        chosen=tuple(c for c in self.cells if r0<=c.row+1<=r1 and c0<=c.column+1<=c1)
        if len(chosen)!=(c1-c0+1)*(r1-r0+1): raise MaterialError('selection_contains_unextracted_cells')
        return chosen


def datasets_from_bundle(extraction_id,bundle,policy=None):
    """Every usable cell is backed by an actual immutable block, not a fake bbox."""
    from openpyxl.utils.cell import get_column_letter
    policy=policy or DatasetPolicy()
    blocks={b.block_id:b for b in bundle.blocks}; groups={}
    for table in (b for b in bundle.blocks if b.kind=='table'):
        group=table.metadata.get('dataset_group',table.block_id)
        groups.setdefault(group,[]).append(table)
    datasets=[]
    for group,tables in groups.items():
        cells=[]; refs=[]; header_names={}; limitations=list(bundle.manifest.limitations)
        effective_columns={c.index:c for c in policy.columns}
        for table in tables:
            for original in table.metadata.get('cells',[]):
                if original['row']==policy.header_row:
                    header_names[original['column']]=original.get('raw',original.get('text','')) or ''
        for col,name in header_names.items():
            if effective_columns.get(col,ColumnPolicy(col)).unit=='auto':
                unit=re.search(r'[([]\s*([^()\[\]]+)\s*[)\]]\s*$',name)
                if unit:
                    try: effective_columns[col]=replace(effective_columns.get(col,ColumnPolicy(col)),unit=unit_code(unit.group(1).strip()))
                    except MaterialError: limitations.append('unknown_header_unit:'+str(col))
        effective_policy=DatasetPolicy(policy.header_row,tuple(effective_columns[c] for c in sorted(effective_columns)))
        for table in tables:
            refs.append(EvidenceRef(bundle.asset_id,bundle.asset_version,extraction_id,table.block_id,table.locator))
            for original in table.metadata.get('cells',[]):
                block=blocks.get(original.get('block_id'))
                if block is None: limitations.append('table_cell_without_evidence'); continue
                row,col=original['row'],original['column']; address=original.get('address',get_column_letter(col+1)+str(row+1))
                raw=original.get('raw',original.get('text','')) or ''
                cp=effective_policy.column(col)
                if original.get('merge')=='continue': kind='merged'
                else: kind=original.get('native_kind')
                value=normalize(raw,cp,native_kind=kind,iso_value=original.get('iso_value'))
                number_format=original.get('number_format') or ''
                if kind=='number' and value.kind=='number':
                    code=re.search(r'\b(RUB|USD|EUR|GBP|CAD|AUD|CNY|JPY)\b',number_format)
                    if code and cp.unit not in ('auto','1',code.group(1)):
                        value=Value('invalid',unit=value.unit,notes=('column_unit_conflict',))
                    elif cp.unit=='auto' and code:
                        value=replace(value,unit=code.group(1))
                    elif cp.unit=='auto' and ('$' in number_format or '£' in number_format):
                        value=Value('ambiguous',notes=('currency_code_required',))
                    elif ('$' in number_format and cp.unit not in ('USD','CAD','AUD')) or ('£' in number_format and cp.unit!='GBP'):
                        value=Value('invalid',unit=value.unit,notes=('column_unit_conflict',))
                    if value.kind=='number' and percent_format(number_format) and cp.unit not in ('auto','1','percent'):
                        value=Value('invalid',unit=value.unit,notes=('column_unit_conflict',))
                    elif value.kind=='number' and percent_format(number_format) and cp.unit in ('auto','percent'):
                        with localcontext() as context:
                            context.prec=100
                            value=replace(value,value=decstr(decimal(value.value)*100),unit='percent',notes=value.notes+('xlsx_percent_fraction_normalized',))
                if row==policy.header_row:
                    header_names[col]=raw
                formula=original.get('formula')
                source=EvidenceRef(bundle.asset_id,bundle.asset_version,extraction_id,block.block_id,block.locator)
                cells.append(DataCell(address,row,col,raw,value,source,block.quality,formula,original.get('cached_raw'),
                    {k:v for k,v in original.items() if k not in ('text','raw','formula','cached_raw','block_id')}))
        if not cells: continue
        if len({c.address for c in cells})!=len(cells): raise MaterialError('table_fragment_address_conflict')
        columns=column_summary(cells,effective_policy,header_names)
        if policy.header_row is not None: limitations.append('first_row_header_policy')
        identity=group if tables[0].metadata.get('dataset_group') else [asdict(tables[0].locator),tables[0].metadata.get('xml_path','')]
        series=sha256(canonical([bundle.asset_id,bundle.extractor,identity,asdict(effective_policy)]).encode()).hexdigest()
        name=str(group) if tables[0].metadata.get('dataset_group') else 'Таблица '+str(len(datasets)+1)+(f' (стр. {tables[0].locator.page})' if tables[0].locator.page else '')
        datasets.append(Dataset(series,name,effective_policy,tuple(sorted(cells,key=lambda c:(c.row,c.column))),tuple(columns),tuple(refs),bundle.manifest.coverage,tuple(dict.fromkeys(limitations))))
    return tuple(datasets)


def column_summary(cells,policy,names):
    from openpyxl.utils.cell import get_column_letter
    columns=[]
    for col in sorted({c.column for c in cells}):
        values=[c.normalized for c in cells if c.column==col and c.row!=policy.header_row]
        kinds={v.kind for v in values if v.kind!='missing'}
        columns.append(dict(index=col,name=names.get(col,get_column_letter(col+1)),
            dtype=next(iter(kinds)) if len(kinds)==1 else 'mixed',
            units=sorted({v.unit for v in values if v.kind=='number'}),policy=asdict(policy.column(col))))
    return tuple(columns)


def correction_proposal(dataset,address,replacement):
    cell=dataset.cell(address)
    value=normalize(replacement,dataset.policy.column(cell.column))
    return dict(dataset_id=dataset.id,address=cell.address,original_raw=cell.raw,proposed_raw=replacement,
        proposed_value=asdict(value),source=asdict(cell.source),status='proposed')


def apply_correction(dataset,proposal,*,author_ref):
    # Caller authenticates author and performs head CAS. Domain function keeps raw.
    if proposal.get('dataset_id')!=dataset.id or proposal.get('status')!='proposed': raise MaterialError('stale_correction')
    cell=dataset.cell(proposal['address'])
    if proposal['original_raw']!=cell.raw: raise MaterialError('stale_correction')
    value=normalize(proposal['proposed_raw'],dataset.policy.column(cell.column))
    corrected=replace(cell,normalized=value,formula=None,cached_raw=None,quality='verified',
        metadata={**cell.metadata,'confirmed_override':proposal['proposed_raw'],'original_formula':cell.formula,'confirmed_by':author_ref})
    cells=tuple(corrected if c.address==cell.address else c for c in dataset.cells)
    columns=column_summary(cells,dataset.policy,{c['index']:c['name'] for c in dataset.columns})
    return replace(dataset,cells=cells,columns=columns,previous_id=dataset.id,
        corrections=dataset.corrections+({**proposal,'status':'confirmed','author_ref':author_ref},))
